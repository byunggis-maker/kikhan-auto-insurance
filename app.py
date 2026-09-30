import hashlib
import json
import os
import re
import subprocess
import tempfile
import unicodedata
import warnings
from html import escape
from pathlib import Path

import requests
import streamlit as st
from dotenv import load_dotenv
from insurance_ocr.image_reader import read_policy_image
from insurance_ocr.validator import validate_policy_result
from insurance_ocr.policy_parser import parse_amount_to_won


# =========================================================
# 기본 설정
# =========================================================

PROJECT_ROOT = Path(__file__).resolve().parent
AUTO_INSURANCE_COVERAGES_FILE = (
    PROJECT_ROOT / "data" / "auto_insurance_coverages.json"
)
with AUTO_INSURANCE_COVERAGES_FILE.open("r", encoding="utf-8") as f:
    INSURANCE_COVERAGE_DB = json.load(f)
load_dotenv(PROJECT_ROOT / ".env")

warnings.filterwarnings(
    "ignore",
    message="urllib3 v2 only supports OpenSSL 1.1.1+.*",
    category=Warning,
)

st.set_page_config(
    page_title="기칸의 자동차보험 보상 나침반",
    page_icon="🧭",
    layout="centered",
)


# =========================================================
# 선택 항목
# =========================================================

USER_POSITIONS = [
    "피해자",
    "사고를 낸 운전자 또는 피보험자",
    "쌍방과실 사고 당사자",
    "아직 판단하기 어려움",
]

INSURANCE_TYPE_OPTIONS = [
    "개인용 자동차보험",
    "업무용 자동차보험",
    "영업용 자동차보험",
    "이륜차 자동차보험",
]

FAULT_RATE_OPTIONS = [
    "아직 정해지지 않음",
    "0%",
    "10%",
    "20%",
    "30%",
    "40%",
    "50%",
    "60%",
    "70%",
    "80%",
    "90%",
    "100%",
    "직접 입력",
]

INJURY_GRADE_OPTIONS = [
    "해당 없음",
    "아직 확인되지 않음",
    "1급",
    "2급",
    "3급",
    "4급",
    "5급",
    "6급",
    "7급",
    "8급",
    "9급",
    "10급",
    "11급",
    "12급",
    "13급",
    "14급",
]

# 약관 PDF에서 확인된 담보명과 특약명만 사용합니다.
# PDF에 근거가 없는 임의의 보상한도, 자기부담금 옵션은 제거했습니다.
COVERAGES_BY_TYPE = {
    "개인용 자동차보험": [
        "대인배상Ⅰ",
        "대인배상Ⅱ",
        "대물배상",
        "자기신체사고",
        "자동차상해",
        "무보험자동차에의한상해",
        "자기차량손해",
        "운전자한정 특약",
        "연령한정 특약",
        "기명피보험자 1인한정 특약",
        "부부한정 특약",
        "가족한정 특약",
        "단기 운전자확대 특약",
    ],
    "업무용 자동차보험": [
        "대인배상Ⅰ",
        "대인배상Ⅱ",
        "대물배상",
        "자기신체사고",
        "자동차상해",
        "무보험자동차에의한상해",
        "자기차량손해",
        "단기 운전자확대 특약",
    ],
    "영업용 자동차보험": [
        "대인배상Ⅰ",
        "대인배상Ⅱ",
        "대물배상",
        "자기신체사고",
        "자동차상해",
        "무보험자동차에의한상해",
        "자기차량손해",
    ],
    "이륜차 자동차보험": [
        "대인배상Ⅰ",
        "대인배상Ⅱ",
        "대물배상",
        "자기신체사고",
        "무보험자동차에의한상해",
        "자기차량손해",
    ],
}

COVERAGE_CODES = {
    "대인배상Ⅰ": "d1",
    "대인배상Ⅱ": "d2",
    "대물배상": "d3",
    "자기신체사고": "d4",
    "자동차상해": "d4a",
    "무보험자동차에의한상해": "d5",
    "자기차량손해": "d6",
    "운전자한정 특약": "x1",
    "연령한정 특약": "x2",
    "기명피보험자 1인한정 특약": "x3",
    "부부한정 특약": "x4",
    "가족한정 특약": "x5",
    "단기 운전자확대 특약": "x6",
}

COST_FIELD_CODES = {
    "치료비": "treat_total",
    "수술비": "surgery",
    "입원비": "inpatient",
    "통원치료비": "outpatient",
    "휴업손해": "lost_income",
    "차량수리비": "repair_cost",
    "견인비": "towing_cost",
    "대차료 또는 교통비": "rental_cost",
    "기타 재산손해": "property_damage",
}


# =========================================================
# 화면 디자인
# =========================================================

st.markdown(
    """
    <style>
    html, body, [data-testid="stAppViewContainer"] {
        background: #ffffff;
        color: #111111;
    }

    .stApp {
        max-width: 1080px;
        margin: 0 auto;
    }

    .block-container {
        padding-top: 2rem;
        padding-bottom: 3rem;
    }

    h1, h2, h3, p, label, .stMarkdown {
        color: #111111;
    }

    div[data-testid="stAlert"] {
        background: #ffffff;
        color: #111111;
        border: 1px solid #333333;
        border-radius: 10px;
    }

    .stTextInput input,
    .stTextArea textarea,
    .stNumberInput input,
    .stSelectbox div[data-baseweb="select"] > div {
        background: #ffffff !important;
        color: #111111 !important;
        border: 1px solid #d0d7de !important;
        border-radius: 8px !important;
        box-shadow: none !important;
    }
    input[type="number"]::-webkit-outer-spin-button,
    input[type="number"]::-webkit-inner-spin-button {
        -webkit-appearance: none;
        margin: 0;
    }
    input[type="number"] {
        -moz-appearance: textfield;
    }
    .white-excel-table {
        width: 100%;
        border-collapse: collapse;
        overflow-x: auto;
        display: block;
        white-space: normal;
        border: 1px solid #dfe3e8;
        border-radius: 10px;
    }
    .white-excel-table th,
    .white-excel-table td {
        border: 1px solid #dfe3e8;
        padding: 0.55rem 0.7rem;
        text-align: left;
        vertical-align: top;
        color: #111111;
        background: #ffffff;
    }
    .white-excel-table th {
        background: #eef3f8;
        font-weight: 700;
    }
    .white-excel-table td.numeric {
        text-align: right;
    }

    .kh-result-card {
        margin: 1rem 0;
        padding: 1.25rem 1.35rem;
        border: 1px solid #dfe6ef;
        border-radius: 16px;
        background: #ffffff;
        box-shadow: 0 8px 24px rgba(29, 54, 92, 0.06);
    }
    .kh-result-label {
        margin-bottom: 0.35rem;
        color: #596579;
        font-size: 0.9rem;
        font-weight: 700;
    }
    .kh-result-amount {
        color: #172033;
        font-size: clamp(1.8rem, 5vw, 2.7rem);
        font-weight: 850;
        letter-spacing: -0.04em;
    }
    .kh-result-note {
        margin-top: 0.55rem;
        color: #596579;
        font-size: 0.92rem;
        line-height: 1.65;
    }
    .kh-action-list {
        margin: 0.7rem 0 0;
        padding-left: 1.2rem;
        color: #263247;
        line-height: 1.75;
    }
    .kh-cost-heading {
        margin: 1rem 0 0.8rem;
        padding: 0.7rem 0.9rem;
        border: 1px solid #e1e5ea;
        border-radius: 8px;
        background: #f2f3f5;
        color: #111111;
        font-weight: 700;
    }
    div[data-testid="stDownloadButton"] > button {
        background: #f2f3f5 !important;
        color: #111111 !important;
        border: 1px solid #d7dce2 !important;
        box-shadow: none !important;
    }
    div[data-testid="stDownloadButton"] > button:hover,
    div[data-testid="stDownloadButton"] > button:focus,
    div[data-testid="stDownloadButton"] > button:active {
        background: #e8eaed !important;
        color: #111111 !important;
        border-color: #c8ced6 !important;
    }
    div[data-testid="stDownloadButton"] > button * {
        color: #111111 !important;
    }

    /* Streamlit expander 제목이 마우스를 떼었을 때 검게 바뀌지 않도록 고정 */
    div[data-testid="stExpander"] details > summary,
    div[data-testid="stExpander"] details > summary:hover,
    div[data-testid="stExpander"] details > summary:focus,
    div[data-testid="stExpander"] details > summary:active {
        background: #f2f3f5 !important;
        color: #111111 !important;
        border: 1px solid #e1e5ea !important;
        border-radius: 8px !important;
        box-shadow: none !important;
    }
    div[data-testid="stExpander"] details > summary * {
        color: #111111 !important;
        fill: #111111 !important;
    }

    @media (max-width: 720px) {
        .block-container {
            padding-left: 1rem !important;
            padding-right: 1rem !important;
        }
        .white-excel-table th,
        .white-excel-table td {
            min-width: 132px;
            padding: 0.5rem 0.58rem;
            font-size: 0.9rem;
        }
    }

    div[data-testid="stAudioInput"] {
        background: #ffffff !important;
        border: 1px solid #111111;
        border-radius: 12px;
        padding: 0.55rem 0.7rem;
    }

    div[data-testid="stAudioInput"] button {
        background: #ffffff !important;
        color: #111111 !important;
        border: 1px solid #111111 !important;
        border-radius: 50% !important;
    }

    div[data-testid="stAudioInput"] button svg {
        color: #111111 !important;
        fill: #111111 !important;
    }

    div.stButton > button,
    div[data-testid="stFormSubmitButton"] > button {
        width: 100%;
        background: #111111 !important;
        color: #ffffff !important;
        border: 1px solid #111111 !important;
        border-radius: 10px;
        font-weight: 700;
        padding: 0.72rem 1rem;
    }

    div.stButton > button:hover,
    div[data-testid="stFormSubmitButton"] > button:hover {
        background: #000000 !important;
        color: #ffffff !important;
    }

    div.stButton > button p,
    div[data-testid="stFormSubmitButton"] > button p {
        color: #ffffff !important;
    }
    
    /* Streamlit 상단 검은 영역 제거 */
    header[data-testid="stHeader"] {
        display: none !important;
        height: 0 !important;
    }

    div[data-testid="stToolbar"] {
        display: none !important;
    }

    div[data-testid="stDecoration"] {
        display: none !important;
    }

    .block-container {
        padding-top: 1.2rem !important;
    }
</style>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# 공통 함수
# =========================================================


def make_storage_key(value):
    return re.sub(
        r"[^0-9A-Za-z가-힣]+",
        "_",
        str(value),
    ).strip("_").lower()


def get_amount_options(coverage_name):
    return [
        "직접 입력",
    ]


def get_deductible_options(coverage_name):
    return [
        "직접 입력",
    ]


def get_openai_key():
    try:
        secret_key = str(
            st.secrets.get(
                "OPENAI_API_KEY",
                "",
            )
        ).strip()
    except Exception:
        secret_key = ""

    return (
        secret_key
        or os.getenv(
            "OPENAI_API_KEY",
            "",
        ).strip()
        or None
    )


def safe_audio_metadata(recorded_audio):
    original_name = str(
        getattr(
            recorded_audio,
            "name",
            "",
        )
        or ""
    ).strip()

    original_type = str(
        getattr(
            recorded_audio,
            "type",
            "",
        )
        or ""
    ).strip()

    mime_type = (
        original_type.split(
            ";",
            1,
        )[0].lower()
        or "audio/wav"
    )

    extension_map = {
        "audio/wav": ".wav",
        "audio/x-wav": ".wav",
        "audio/mpeg": ".mp3",
        "audio/mp3": ".mp3",
        "audio/webm": ".webm",
        "audio/mp4": ".m4a",
        "audio/x-m4a": ".m4a",
        "audio/ogg": ".ogg",
    }

    suffix = Path(
        original_name
    ).suffix.lower()

    if suffix not in {
        ".wav",
        ".mp3",
        ".webm",
        ".m4a",
        ".ogg",
        ".mp4",
    }:
        suffix = extension_map.get(
            mime_type,
            ".wav",
        )

    return (
        f"accident_audio{suffix}",
        mime_type,
    )


def normalize_audio_to_wav(audio_bytes, mime_type="audio/wav"):
    if not audio_bytes:
        raise RuntimeError("녹음된 음성 데이터가 비어 있습니다.")

    if audio_bytes.startswith(b"RIFF"):
        return audio_bytes, "audio/wav"

    if mime_type not in {"audio/webm", "audio/ogg", "audio/m4a", "audio/mp4", "audio/mpeg", "audio/mp3"}:
        return audio_bytes, mime_type

    with tempfile.NamedTemporaryFile(suffix=".input") as src, tempfile.NamedTemporaryFile(suffix=".wav") as dst:
        src.write(audio_bytes)
        src.flush()

        cmd = [
            "ffmpeg",
            "-y",
            "-i",
            src.name,
            "-vn",
            "-ac",
            "1",
            "-ar",
            "16000",
            "-f",
            "wav",
            dst.name,
        ]
        result = subprocess.run(
            cmd,
            capture_output=True,
            check=False,
        )

        if result.returncode != 0:
            stderr = (result.stderr or b"").decode("utf-8", errors="replace")
            raise RuntimeError(
                "브라우저 녹음 파일을 변환하지 못했습니다. "
                "다시 녹음하거나 WAV 형식으로 저장한 뒤 다시 시도해 주세요."
                f"{(' ' + stderr[:200]) if stderr else ''}"
            )

        wav_bytes = dst.read()

    return wav_bytes, "audio/wav"


def transcribe_audio(
    audio_bytes,
    file_name,
    mime_type,
):
    api_key = get_openai_key()

    if not api_key:
        raise RuntimeError(
            "OPENAI_API_KEY가 설정되지 않았습니다. "
            "로컬은 .env, 공개 앱은 "
            "Streamlit Secrets를 확인해 주세요."
        )

    normalized_audio, normalized_mime = normalize_audio_to_wav(
        audio_bytes,
        mime_type,
    )
    safe_file_name = Path(file_name).stem + ".wav"

    if (
        not normalized_audio
        or len(normalized_audio) < 800
    ):
        raise RuntimeError(
            "녹음 내용이 너무 짧습니다. "
            "2초 이상 말한 뒤 녹음을 종료해 주세요."
        )

    try:
        response = requests.post(
            "https://api.openai.com/v1/audio/transcriptions",
            headers={
                "Authorization": (
                    f"Bearer {api_key}"
                )
            },
            files={
                "file": (
                    safe_file_name,
                    normalized_audio,
                    normalized_mime,
                )
            },
            data={
                "model": "gpt-4o-mini-transcribe",
                "language": "ko",
                "response_format": "json",
            },
            timeout=(15, 150),
        )

    except requests.exceptions.Timeout as exc:
        raise RuntimeError(
            "변환 서버 응답이 지연되었습니다. "
            "같은 녹음으로 다시 시도해 주세요."
        ) from exc

    except requests.exceptions.RequestException as exc:
        raise RuntimeError(
            "음성 변환 서버에 연결하지 못했습니다. "
            "인터넷 연결을 확인해 주세요."
        ) from exc

    if not response.ok:
        try:
            error_data = response.json().get(
                "error",
                {},
            )

            error_message = str(
                error_data.get(
                    "message",
                    "",
                )
            ).strip()

        except Exception:
            error_message = response.text[:400]

        safe_message = re.sub(
            r"sk-[A-Za-z0-9_-]{10,}",
            "[API_KEY_REDACTED]",
            error_message,
        ).replace(
            "\n",
            " ",
        ).strip()

        lowered = safe_message.lower()

        if (
            response.status_code == 401
            or "invalid_api_key" in lowered
        ):
            raise RuntimeError(
                "API 키 인증에 실패했습니다. "
                "키 저장 위치를 확인해 주세요."
            )

        if response.status_code == 429:
            raise RuntimeError(
                "API 사용 한도 또는 "
                "요청 한도를 확인해 주세요."
            )

        if response.status_code == 413:
            raise RuntimeError(
                "녹음 파일이 너무 큽니다. "
                "더 짧게 녹음해 주세요."
            )

        raise RuntimeError(
            "음성 변환 요청이 거절되었습니다. "
            f"상태번호 {response.status_code}: "
            f"{safe_message or '상세 원인을 확인하지 못했습니다.'}"
        )

    try:
        result = response.json()

    except ValueError as exc:
        raise RuntimeError(
            "변환 결과를 읽지 못했습니다. "
            "같은 녹음으로 다시 시도해 주세요."
        ) from exc

    transcript = str(
        result.get(
            "text",
            "",
        )
    ).strip()

    if not transcript:
        raise RuntimeError(
            "말소리를 확인하지 못했습니다. "
            "마이크에 가까이 대고 다시 녹음해 주세요."
        )

    return transcript


def infer_related_coverages(selected_coverages, question_text):
    question = (question_text or "").lower()
    if not selected_coverages:
        return ["관련 담보 확인 필요"]

    selected_names = [item.get("담보명") for item in selected_coverages]
    injury_keywords = ["치료", "부상", "입원", "통원", "수술", "후유장해", "휴업", "소득", "상해"]
    vehicle_keywords = ["수리", "차량", "견인", "대차", "교통", "손상", "파손"]
    property_keywords = ["물건", "대물", "파손", "창문", "타이어"]

    if any(keyword in question for keyword in injury_keywords):
        candidates = [name for name in selected_names if name in {"대인배상Ⅰ", "대인배상Ⅱ", "자기신체사고", "자동차상해", "무보험자동차에의한상해"}]
        return candidates or ["관련 담보 확인 필요"]

    if any(keyword in question for keyword in vehicle_keywords):
        candidates = [name for name in selected_names if name == "자기차량손해"]
        if candidates:
            return candidates
        return ["관련 담보 확인 필요"]

    if any(keyword in question for keyword in property_keywords):
        candidates = [name for name in selected_names if name == "대물배상"]
        if candidates:
            return candidates
        return ["관련 담보 확인 필요"]

    return selected_names[:1] if selected_names else ["관련 담보 확인 필요"]


def normalize_policy_search_text(value):
    return re.sub(
        r"[^0-9A-Za-z가-힣ⅠⅡ]+",
        "",
        str(value or ""),
    ).lower()


POLICY_COVERAGE_ALIASES = {
    "대인배상Ⅰ": [
        "대인배상Ⅰ",
        "대인배상1",
        "대인1",
        "대인 일",
    ],
    "대인배상Ⅱ": [
        "대인배상Ⅱ",
        "대인배상2",
        "대인2",
        "대인 이",
    ],
    "대물배상": [
        "대물배상",
        "대물",
        "차량수리",
        "상대 차량 손해",
    ],
    "자기신체사고": [
        "자기신체사고",
        "자손",
        "본인 치료비",
    ],
    "자동차상해": [
        "자동차상해",
        "자동차상해특약",
        "자상",
    ],
    "무보험자동차에의한상해": [
        "무보험자동차에의한상해",
        "무보험자동차상해",
        "무보험차상해",
        "무보험차 상해",
        "무보험 차량 상해",
        "무보험",
    ],
    "자기차량손해": [
        "자기차량손해",
        "자차",
        "자차보험",
        "내 차 수리",
    ],
    "운전자한정 특약": [
        "운전자한정",
        "운전자 한정",
        "운전자 범위",
    ],
    "연령한정 특약": [
        "연령한정",
        "연령 한정",
        "운전자 나이",
    ],
}


def detect_requested_coverage(question_text):
    normalized_question = normalize_policy_search_text(
        question_text
    )

    matches = []

    for coverage_name, aliases in POLICY_COVERAGE_ALIASES.items():
        for alias in aliases:
            normalized_alias = normalize_policy_search_text(
                alias
            )

            if (
                normalized_alias
                and normalized_alias in normalized_question
            ):
                matches.append(
                    (
                        len(normalized_alias),
                        coverage_name,
                    )
                )

    if not matches:
        return None

    matches.sort(
        reverse=True
    )

    return matches[0][1]


def get_question_search_terms(question_text, coverage_name=None):
    terms = []

    if coverage_name:
        terms.append(
            coverage_name
        )

        terms.extend(
            POLICY_COVERAGE_ALIASES.get(
                coverage_name,
                [],
            )
        )

    stop_words = {
        "약관",
        "규정",
        "내용",
        "알려줘",
        "알려주세요",
        "설명",
        "설명해줘",
        "어떻게",
        "경우",
        "관련",
        "대한",
        "좀",
    }

    for token in re.findall(
        r"[0-9A-Za-z가-힣ⅠⅡ]{2,}",
        str(question_text or ""),
    ):
        if token not in stop_words:
            terms.append(token)

    unique_terms = []
    seen = set()

    for term in terms:
        normalized = normalize_policy_search_text(
            term
        )

        if not normalized or normalized in seen:
            continue

        seen.add(normalized)
        unique_terms.append(term)

    return unique_terms


def normalize_extracted_pdf_text(raw_text):
    if not raw_text:
        return ""

    lines = [
        line.rstrip()
        for line in str(raw_text).splitlines()
    ]

    cleaned_lines = []
    vertical_fragments = []

    def is_vertical_fragment(value):
        stripped = value.strip()

        if not stripped:
            return False

        return bool(
            re.fullmatch(
                r"[0-9A-Za-z가-힣ⅠⅡ]{1,2}",
                stripped,
            )
        )

    def flush_vertical_fragments():
        nonlocal vertical_fragments

        if not vertical_fragments:
            return

        if len(vertical_fragments) >= 3:
            joined = "".join(
                fragment.strip()
                for fragment in vertical_fragments
            )

            cleaned_lines.append(
                joined
            )
        else:
            cleaned_lines.extend(
                vertical_fragments
            )

        vertical_fragments = []

    for line in lines:
        stripped = line.strip()

        if is_vertical_fragment(
            stripped
        ):
            vertical_fragments.append(
                stripped
            )
            continue

        flush_vertical_fragments()

        if not stripped:
            if (
                cleaned_lines
                and cleaned_lines[-1] != ""
            ):
                cleaned_lines.append("")
            continue

        cleaned_lines.append(
            stripped
        )

    flush_vertical_fragments()

    result = "\n".join(
        cleaned_lines
    )

    result = re.sub(
        r"\n{3,}",
        "\n\n",
        result,
    )

    return result.strip()


def find_policy_original_in_pdfs(
    question_text,
    insurance_type=None,
):
    try:
        from pypdf import PdfReader
    except ImportError:
        return {
            "원문": "",
            "PDF": "",
            "페이지": "",
            "오류": "PDF 검색 기능에 필요한 pypdf가 설치되지 않았습니다.",
        }

    coverage_name = detect_requested_coverage(
        question_text
    )

    search_terms = get_question_search_terms(
        question_text,
        coverage_name,
    )

    normalized_terms = [
        normalize_policy_search_text(term)
        for term in search_terms
        if normalize_policy_search_text(term)
    ]

    policy_pdf_map = {
        "개인용 자동차보험": "개인용 약관.pdf",
        "업무용 자동차보험": "업무용 약관.pdf",
        "영업용 자동차보험": "영업용 약관.pdf",
        "이륜차 자동차보험": "이륜차 약관.pdf",
    }

    target_pdf_name = policy_pdf_map.get(insurance_type)

    if target_pdf_name:
        pdf_paths = sorted(
            path
            for path in PROJECT_ROOT.rglob("*.pdf")
            if unicodedata.normalize("NFC", path.name)
            == unicodedata.normalize("NFC", target_pdf_name)
        )
    else:
        pdf_paths = sorted(
            PROJECT_ROOT.rglob("*.pdf")
        )

    if not pdf_paths:
        return {
            "원문": "",
            "PDF": "",
            "페이지": "",
            "관련 담보": coverage_name or "확인 필요",
            "오류": "프로젝트 폴더에서 자동차보험 약관 PDF를 찾지 못했습니다.",
        }

    insurance_tokens = [
        token
        for token in re.findall(
            r"[가-힣A-Za-z0-9]{2,}",
            str(insurance_type or ""),
        )
        if token not in {
            "자동차보험",
            "보험",
        }
    ]

    normalized_coverage_names = {
        name: normalize_policy_search_text(name)
        for name in POLICY_COVERAGE_ALIASES
    }

    def pdf_priority(pdf_path):
        normalized_name = normalize_policy_search_text(
            pdf_path.name
        )

        score = 0

        for token in insurance_tokens:
            normalized_token = normalize_policy_search_text(
                token
            )
            if normalized_token in normalized_name:
                score += 20

        return score

    pdf_paths.sort(
        key=pdf_priority,
        reverse=True,
    )

    best_candidate = None

    for pdf_path in pdf_paths:
        try:
            reader = PdfReader(
                str(pdf_path)
            )
        except Exception:
            continue

        extracted_pages = []

        for page in reader.pages:
            try:
                try:
                    raw_page_text = page.extract_text(
                        extraction_mode="layout"                    )
                except TypeError:
                    raw_page_text = page.extract_text()

                page_text = normalize_extracted_pdf_text(
                    raw_page_text
                    or ""
                )
            except Exception:
                page_text = ""

            extracted_pages.append(
                page_text
            )

        for page_index, page_text in enumerate(
            extracted_pages
        ):
            if not page_text:
                continue

            normalized_page = normalize_policy_search_text(
                page_text
            )

            score = pdf_priority(
                pdf_path
            )

            coverage_names_on_page = []

            for name, normalized_name in normalized_coverage_names.items():
                if normalized_name and normalized_name in normalized_page:
                    coverage_names_on_page.append(
                        name
                    )

            if coverage_name:
                canonical = normalize_policy_search_text(
                    coverage_name
                )

                aliases = [
                    normalize_policy_search_text(alias)
                    for alias in POLICY_COVERAGE_ALIASES.get(
                        coverage_name,
                        [],
                    )
                ]

                matched_aliases = [
                    alias
                    for alias in aliases
                    if alias and alias in normalized_page
                ]

                if (
                    canonical not in normalized_page
                    and not matched_aliases
                ):
                    continue

                if canonical in normalized_page:
                    score += 150

                score += len(
                    matched_aliases
                ) * 20

                if "보상하는손해" in normalized_page:
                    score += 100

                if "피보험자" in normalized_page:
                    score += 40

                if "보험금" in normalized_page:
                    score += 30

                if "제1조" in normalized_page:
                    score += 50

                if "목차" in normalized_page:
                    score -= 200

                if len(
                    coverage_names_on_page
                ) > 1:
                    score -= (
                        len(
                            coverage_names_on_page
                        )
                        - 1
                    ) * 80

            else:
                term_hits = 0

                for term in normalized_terms:
                    if term in normalized_page:
                        term_hits += 1

                if term_hits == 0:
                    continue

                score += term_hits * 10

            candidate = {
                "점수": score,
                "PDF경로": pdf_path,
                "PDF명": pdf_path.name,
                "시작페이지": page_index,
                "전체페이지": extracted_pages,
            }

            if (
                best_candidate is None
                or candidate["점수"]
                > best_candidate["점수"]
            ):
                best_candidate = candidate

    if not best_candidate:
        return {
            "원문": "",
            "PDF": "",
            "페이지": "",
            "관련 담보": coverage_name or "확인 필요",
            "오류": (
                "등록된 약관 PDF에서 질문과 일치하는 "
                "약관 조항을 찾지 못했습니다."
            ),
        }

    start_page = best_candidate[
        "시작페이지"
    ]

    all_pages = best_candidate[
        "전체페이지"
    ]

    collected_pages = []

    max_end = min(
        len(all_pages),
        start_page + 15,
    )

    for page_index in range(
        start_page,
        max_end,
    ):
        page_text = all_pages[
            page_index
        ]

        if not page_text:
            continue

        normalized_page = normalize_policy_search_text(
            page_text
        )

        if (
            page_index > start_page
            and coverage_name
        ):
            beginning = normalized_page[
                :700
            ]

            next_coverage_found = False

            for other_name, normalized_name in normalized_coverage_names.items():
                if other_name == coverage_name:
                    continue

                if (
                    normalized_name
                    and normalized_name in beginning
                    and (
                        "보상하는손해" in beginning
                        or "제1조" in beginning
                        or "제1절" in beginning
                    )
                ):
                    next_coverage_found = True
                    break

            if next_coverage_found:
                break

        collected_pages.append(
            (
                page_index + 1,
                page_text,
            )
        )

    if not collected_pages:
        return {
            "원문": "",
            "PDF": best_candidate[
                "PDF명"
            ],
            "페이지": start_page + 1,
            "관련 담보": coverage_name or "확인 필요",
            "오류": "관련 페이지에서 읽을 수 있는 약관 원문을 찾지 못했습니다.",
        }

    full_original_text = []

    for page_number, page_text in collected_pages:
        full_original_text.append(
            f"[PDF {page_number}쪽]\n\n{page_text}"
        )

    first_page = collected_pages[0][0]
    last_page = collected_pages[-1][0]

    if first_page == last_page:
        page_display = str(
            first_page
        )
    else:
        page_display = (
            f"{first_page}~{last_page}"
        )

    return {
        "관련 담보": coverage_name or "질문 내용 기준 검색",
        "원문": "\n\n━━━━━━━━━━━━━━━━━━━━\n\n".join(
            full_original_text
        ),
        "PDF": best_candidate[
            "PDF명"
        ],
        "페이지": page_display,
        "오류": "",
    }

def find_policy_original_in_rules(
    question_text,
    insurance_type=None,
):
    coverage_name = detect_requested_coverage(
        question_text
    )

    candidates = []

    for rule in POLICY_RULES:
        if (
            insurance_type
            and rule.get("보험 종류")
            and rule.get("보험 종류") != insurance_type
        ):
            continue

        rule_coverage = str(
            rule.get(
                "담보명",
                "",
            )
        )

        if (
            coverage_name
            and rule_coverage != coverage_name
        ):
            continue

        searchable = " ".join(
            str(value)
            for value in rule.values()
            if value is not None
        )

        score = 0

        for term in get_question_search_terms(
            question_text,
            coverage_name,
        ):
            if (
                normalize_policy_search_text(term)
                in normalize_policy_search_text(searchable)
            ):
                score += 1

        if coverage_name:
            score += 20

        if score > 0:
            candidates.append(
                (
                    score,
                    rule,
                )
            )

    if not candidates:
        return None

    candidates.sort(
        key=lambda item: item[0],
        reverse=True,
    )

    rule = candidates[0][1]

    original_text = ""

    for key in [
        "약관 원문",
        "조문 원문",
        "조항 원문",
        "원문",
        "본문",
    ]:
        value = str(
            rule.get(
                key,
                "",
            )
            or ""
        ).strip()

        if value:
            original_text = value
            break

    if not original_text:
        return {
            "관련 담보": rule.get(
                "담보명",
                coverage_name or "확인 필요",
            ),
            "원문": "",
            "PDF": rule.get(
                "PDF 파일명",
                "",
            ),
            "페이지": rule.get(
                "PDF 페이지",
                "",
            ),
            "오류": (
                "약관 규칙 자료에는 요약값만 있고 "
                "PDF 원문이 저장되어 있지 않습니다."
            ),
        }

    return {
        "관련 담보": rule.get(
            "담보명",
            coverage_name or "확인 필요",
        ),
        "원문": original_text,
        "PDF": rule.get(
            "PDF 파일명",
            "",
        ),
        "페이지": rule.get(
            "PDF 페이지",
            "",
        ),
        "오류": "",
    }


def build_consultation_answer(
    selected_coverages,
    question_text,
    insurance_type=None,
):
    pdf_result = find_policy_original_in_pdfs(
        question_text,
        insurance_type,
    )

    if pdf_result.get(
        "원문"
    ):
        return pdf_result

    rule_result = find_policy_original_in_rules(
        question_text,
        insurance_type,
    )

    if (
        rule_result
        and rule_result.get(
            "원문"
        )
    ):
        return rule_result

    if rule_result:
        return rule_result

    return pdf_result


# 보호 구역: 실제 음성변환 회귀시험 없이 이 코드를 수정하지 말 것
# 이 섹션은 마이크 입력, OpenAI 변환, 세션 캐시와 rerun 보호 로직만 담당한다.
def initialize_state():
    defaults = {
        "accident_input_text": "",
        "processed_audio_hash": "",
        "failed_audio_hash": "",
        "stt_error": "",
        "force_stt_retry": False,
        "transcription_cache": {},
        "analysis_result": None,
        "policy_question_text": "",
        "question_processed_audio_hash": "",
        "question_failed_audio_hash": "",
        "question_stt_error": "",
        "question_force_stt_retry": False,
        "policy_question_answer": None,
        "policy_ocr_result": None,
        "policy_ocr_error": "",
        "policy_ocr_image_hash": "",
        "policy_ocr_coverage_meta": {},
        "insurance_type_select": INSURANCE_TYPE_OPTIONS[0],
    }

    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value


def process_recorded_audio(recorded_audio):
    try:
        audio_bytes = recorded_audio.getvalue()
    except Exception as exc:
        st.session_state.stt_error = (
            "녹음된 오디오를 읽지 못했습니다. 마이크 권한과 브라우저 녹음 상태를 다시 확인해 주세요."
        )
        st.session_state.failed_audio_hash = ""
        return

    if not audio_bytes:
        st.session_state.stt_error = (
            "녹음된 파일이 비어 있습니다. 마이크에 가까이 대고 5초 이상 말한 뒤 다시 녹음해 주세요."
        )
        st.session_state.failed_audio_hash = ""
        return

    audio_hash = hashlib.sha256(
        audio_bytes
    ).hexdigest()

    cached_text = st.session_state.get("transcription_cache", {}).get(audio_hash)
    if cached_text:
        st.session_state.accident_input_text = cached_text
        st.session_state.processed_audio_hash = audio_hash
        st.session_state.failed_audio_hash = ""
        st.session_state.stt_error = ""
        return

    should_retry = bool(
        st.session_state.get(
            "force_stt_retry"
        )
    )

    already_succeeded = (
        audio_hash
        == st.session_state.get(
            "processed_audio_hash"
        )
    )

    already_failed = (
        audio_hash
        == st.session_state.get(
            "failed_audio_hash"
        )
    )

    if already_succeeded:
        return

    if (
        already_failed
        and not should_retry
    ):
        return

    st.session_state.force_stt_retry = False

    file_name, mime_type = safe_audio_metadata(
        recorded_audio
    )

    try:
        with st.spinner(
            "음성을 문자로 바꾸는 중입니다"
        ):
            transcript = transcribe_audio(
                audio_bytes,
                file_name,
                mime_type,
            )

        st.session_state.accident_input_text = (
            transcript
        )
        st.session_state.transcription_cache[audio_hash] = transcript
        st.session_state.processed_audio_hash = (
            audio_hash
        )
        st.session_state.failed_audio_hash = ""
        st.session_state.stt_error = ""

    except Exception as error:
        st.session_state.failed_audio_hash = (
            audio_hash
        )
        st.session_state.stt_error = str(
            error
        )



def process_question_recorded_audio(recorded_audio):
    try:
        audio_bytes = recorded_audio.getvalue()
    except Exception:
        st.session_state.question_stt_error = (
            "질문 녹음파일을 읽지 못했습니다. "
            "Chrome의 마이크 권한을 확인해 주세요."
        )
        return

    if not audio_bytes:
        st.session_state.question_stt_error = (
            "질문 녹음 내용이 비어 있습니다. "
            "2초 이상 말한 뒤 녹음을 종료해 주세요."
        )
        return

    audio_hash = hashlib.sha256(audio_bytes).hexdigest()

    cached_text = st.session_state.get(
        "transcription_cache",
        {},
    ).get(audio_hash)

    if cached_text:
        st.session_state.policy_question_text = cached_text
        st.session_state.question_processed_audio_hash = audio_hash
        st.session_state.question_failed_audio_hash = ""
        st.session_state.question_stt_error = ""
        return

    already_succeeded = (
        audio_hash
        == st.session_state.get(
            "question_processed_audio_hash"
        )
    )

    already_failed = (
        audio_hash
        == st.session_state.get(
            "question_failed_audio_hash"
        )
    )

    should_retry = bool(
        st.session_state.get(
            "question_force_stt_retry"
        )
    )

    if already_succeeded:
        return

    if already_failed and not should_retry:
        return

    st.session_state.question_force_stt_retry = False

    file_name, mime_type = safe_audio_metadata(
        recorded_audio
    )

    try:
        with st.spinner(
            "약관 질문을 문자로 바꾸는 중입니다"
        ):
            transcript = transcribe_audio(
                audio_bytes,
                file_name,
                mime_type,
            )

        st.session_state.policy_question_text = transcript
        st.session_state.transcription_cache[audio_hash] = transcript
        st.session_state.question_processed_audio_hash = audio_hash
        st.session_state.question_failed_audio_hash = ""
        st.session_state.question_stt_error = ""
        st.session_state.policy_question_answer = None

    except Exception as error:
        st.session_state.question_failed_audio_hash = audio_hash
        st.session_state.question_stt_error = str(error)


def reflow_policy_text_for_reading(original_text):
    if not original_text:
        return ""

    text_value = str(
        original_text
    ).replace(
        "\u200b",
        "",
    ).replace(
        "\ufeff",
        "",
    )

    lines = text_value.splitlines()
    paragraphs = []
    sentence_buffer = []

    def flush_sentence_buffer():
        nonlocal sentence_buffer

        if sentence_buffer:
            paragraph = " ".join(
                sentence_buffer
            )

            paragraph = re.sub(
                r"[ \t]+",
                " ",
                paragraph,
            ).strip()

            if paragraph:
                paragraphs.append(
                    paragraph
                )

        sentence_buffer = []

    for raw_line in lines:
        line = re.sub(
            r"[ \t]+",
            " ",
            raw_line,
        ).strip()

        if not line:
            flush_sentence_buffer()
            continue

        if set(line) <= {"━", "-", "─", "—"}:
            flush_sentence_buffer()
            continue

        if re.fullmatch(
            r"\[PDF\s*\d+쪽\]",
            line,
        ):
            flush_sentence_buffer()
            paragraphs.append(
                line
            )
            continue

        is_clause_heading = bool(
            re.match(
                r"^(제\s*\d+\s*[조장절편]|"
                r"[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]|"
                r"\d+\.\s*|"
                r"[가-힣]\.\s*)",
                line,
            )
        )

        is_short_heading = (
            len(line) <= 32
            and not re.search(
                r"[.!?。]$",
                line,
            )
            and any(
                word in line
                for word in [
                    "보상하는 손해",
                    "보상하지 않는 손해",
                    "피보험자",
                    "보험금",
                    "지급기준",
                    "용어의 정의",
                    "청구",
                    "면책",
                    "담보",
                    "특약",
                ]
            )
        )

        if (
            is_clause_heading
            or is_short_heading
        ):
            flush_sentence_buffer()
            sentence_buffer.append(
                line
            )

            if is_short_heading:
                flush_sentence_buffer()

            continue

        sentence_buffer.append(
            line
        )

        if re.search(
            r"(다\.|니다\.|습니다\.|[.!?。])$",
            line,
        ):
            flush_sentence_buffer()

    flush_sentence_buffer()

    return "\n\n".join(
        paragraphs
    )


def render_policy_question_section(result):
    st.subheader("약관 내용 확인")

    st.caption(
        "원하는 담보의 약관 규정을 말하거나 직접 입력해 주세요. "
        "등록된 자동차보험 약관 PDF에서 일치하는 원문을 찾아 "
        "요약하지 않고 그대로 표시합니다."
    )

    question_audio = st.audio_input(
        "약관 질문 녹음 시작",
        key="policy_question_audio",
        label_visibility="visible",
    )

    if question_audio is not None:
        process_question_recorded_audio(
            question_audio
        )

    if st.session_state.get(
        "question_stt_error"
    ):
        st.error(
            "질문을 문자로 바꾸지 못했습니다. "
            f"{st.session_state.question_stt_error}"
        )

        if st.button(
            "같은 질문 녹음 다시 변환",
            key="retry_policy_question_stt",
        ):
            st.session_state.question_force_stt_retry = True
            st.session_state.question_failed_audio_hash = ""
            st.session_state.question_stt_error = ""
            st.rerun()

    elif st.session_state.get(
        "question_processed_audio_hash"
    ):
        st.success(
            "질문이 문자로 변환되었습니다. "
            "아래 문장을 확인하거나 수정해 주세요."
        )

    question_text = st.text_area(
        "질문 내용 확인 및 수정",
        key="policy_question_text",
        height=130,
        placeholder=(
            "예: 무보험자동차에 의한 상해 담보의 "
            "약관 규정을 보여주세요."
        ),
    )

    if st.button(
        "약관 내용 찾기",
        key="find_policy_answer",
    ):
        if not question_text.strip():
            st.warning(
                "찾으려는 담보나 약관 내용을 "
                "말하거나 직접 입력해 주세요."
            )
        else:
            st.session_state.policy_question_answer = (
                build_consultation_answer(
                    result.get(
                        "selected_coverages",
                        [],
                    ),
                    question_text,
                    result.get(
                        "insurance_type",
                    ),
                )
            )

    answer = st.session_state.get(
        "policy_question_answer"
    )

    if not answer:
        return

    error_message = str(
        answer.get(
            "오류",
            "",
        )
        or ""
    ).strip()

    original_text = str(
        answer.get(
            "원문",
            "",
        )
        or ""
    ).strip()

    if error_message and not original_text:
        st.warning(
            error_message
        )
        return

    st.markdown("#### 관련 약관 규정")

    st.markdown(
        "##### PDF 약관 원문"
    )

    reading_text = reflow_policy_text_for_reading(
        original_text
    )

    st.markdown(
        "<div style=\""
        "display:block;"
        "width:100%;"
        "max-width:100%;"
        "writing-mode:horizontal-tb !important;"
        "text-orientation:mixed !important;"
        "white-space:pre-wrap;"
        "word-break:normal;"
        "overflow-wrap:break-word;"
        "max-height:620px;"
        "overflow-y:auto;"
        "overflow-x:hidden;"
        "padding:24px;"
        "box-sizing:border-box;"
        "background:#ffffff;"
        "color:#111111;"
        "border:1px solid #d0d7de;"
        "border-radius:10px;"
        "line-height:1.85;"
        "font-size:16px;"
        "font-family:-apple-system,BlinkMacSystemFont,"
        "'Apple SD Gothic Neo','Noto Sans KR',sans-serif;"
        "text-align:left;"
        "\">"
        + escape(reading_text)
        + "</div>",
        unsafe_allow_html=True,
    )

    page_value = answer.get(
        "페이지",
        "",
    )

    if page_value:
        st.markdown(
            f"**관련 페이지: {page_value}쪽**"
        )
    else:
        st.markdown(
            "**관련 페이지: 확인 필요**"
        )

    pdf_name = str(
        answer.get(
            "PDF",
            "",
        )
        or ""
    ).strip()

    if pdf_name:
        st.caption(
            f"출처: {pdf_name}"
        )

def format_currency_value(raw_value):
    if raw_value is None:
        return "미입력"

    value = str(raw_value).strip()
    if value == "":
        return "미입력"

    if value.lower() in {
        "미입력",
        "선택하세요",
    }:
        return "미입력"

    digits = re.sub(r"[^0-9-]", "", value)
    if digits in {"", "-"}:
        return "미입력"

    try:
        number = int(digits)
    except ValueError:
        return "미입력"

    return f"{number:,.0f}원"


def extract_numeric_digits(raw_value):
    if raw_value is None:
        return ""

    text = str(raw_value).strip()
    if text == "":
        return ""

    if text.lower() in {"미입력", "none", "null", "선택하세요"}:
        return ""

    digits = re.sub(r"[^0-9]", "", text)
    return digits


def format_currency_text(raw_value):
    digits = extract_numeric_digits(raw_value)
    if digits == "":
        return ""
    return f"{int(digits):,}"


def convert_currency_to_int(raw_value):
    digits = extract_numeric_digits(raw_value)
    if digits == "":
        return None
    return int(digits)


def apply_currency_mask(raw_value):
    digits = extract_numeric_digits(raw_value)
    if digits == "":
        return ""
    return f"{int(digits):,}"


def format_currency_widget_value(key):
    st.session_state[key] = apply_currency_mask(
        st.session_state.get(key, "")
    )


def render_money_input(label, key, help_text=None):
    if key not in st.session_state:
        st.session_state[key] = ""
    else:
        st.session_state[key] = apply_currency_mask(
            st.session_state.get(key, "")
        )

    input_value = st.text_input(
        label,
        key=key,
        help=help_text,
        placeholder="0",
        max_chars=18,
        on_change=format_currency_widget_value,
        args=(key,),
    )
    digits = extract_numeric_digits(input_value)
    if input_value and not digits:
        st.caption("숫자만 입력할 수 있습니다. 천 단위 쉼표는 자동으로 표시됩니다.")

    return convert_currency_to_int(input_value)


def render_white_table(headers, rows):
    header_html = "".join(
        f"<th style='white-space:normal; word-break:keep-all; line-height:1.45;'>{escape(str(header))}</th>"
        for header in headers
    )

    body_html = []
    for row in rows:
        cells = "".join(
            "<td style='white-space:normal; word-break:keep-all; "
            "overflow-wrap:anywhere; line-height:1.55; vertical-align:top;'>"
            f"{escape(str(row.get(header, '')))}</td>"
            for header in headers
        )

        is_total = (
            str(row.get("가입담보명", "")).strip() == "합계"
            or "예상 보험금 합계" in str(row.get("담보 구분", ""))
        )
        row_style = (
            " style='font-weight:700; background:#f3f6f9;'"
            if is_total
            else ""
        )
        body_html.append(f"<tr{row_style}>{cells}</tr>")

    st.markdown(
        "<div style='overflow-x:auto; width:100%;'>"
        "<table class='white-excel-table'>"
        f"<thead><tr>{header_html}</tr></thead>"
        f"<tbody>{''.join(body_html)}</tbody>"
        "</table></div>",
        unsafe_allow_html=True,
    )


def build_analysis_summary(
    position,
    accident_text,
    insurance_type,
    insurer,
    insurance_product,
    insurance_period,
    injury_grade,
    disability_grade,
    hospital_days,
    fault_rate,
    selected_coverages,
    preconditions=None,
):
    missing = []

    if not accident_text.strip():
        missing.append("사고내용")

    if not insurer.strip():
        missing.append("보험회사")

    if not insurance_period.strip():
        missing.append("보험기간")

    if not selected_coverages:
        missing.append("가입담보")

    return {
        "position": position,
        "accident_text": accident_text.strip(),
        "insurance_type": insurance_type,
        "insurer": insurer.strip(),
        "insurance_product": (insurance_product or "").strip(),
        "insurance_period": insurance_period.strip(),
        "injury_grade": injury_grade,
        "disability_grade": disability_grade,
        "hospital_days": int(hospital_days) if isinstance(hospital_days, (int, float)) and hospital_days >= 0 else 0,
        "fault_rate": fault_rate,
        "selected_coverages": selected_coverages,
        "preconditions": preconditions or {},
        "missing": missing,
    }


def load_policy_rules():
    rule_path = PROJECT_ROOT / "data" / "processed" / "policy_rules.json"
    if not rule_path.exists():
        return []

    try:
        with rule_path.open("r", encoding="utf-8") as source:
            data = json.load(source)
        if isinstance(data, list):
            return data
    except Exception:
        return []

    return []


POLICY_RULES = load_policy_rules()


def get_policy_rule(insurance_type, coverage_name):
    for item in POLICY_RULES:
        if (
            item.get("보험 종류") == insurance_type
            and item.get("담보명") == coverage_name
        ):
            return item

    if coverage_name == "자동차상해":
        auto_injury_sources = {
            "개인용 자동차보험": ("개인용 약관.pdf", "119-120"),
            "업무용 자동차보험": ("업무용 약관.pdf", "84-85"),
            "영업용 자동차보험": ("영업용 약관.pdf", "73-74"),
        }
        source = auto_injury_sources.get(insurance_type)
        if source:
            return {
                "보험 종류": insurance_type,
                "담보명": coverage_name,
                "계산식": "지급보험금 = 실제손해액 + 비용 - 공제액",
                "PDF 파일명": source[0],
                "PDF 페이지": source[1],
                "조문명": "자동차상해 특별약관 2. 보상하는 손해",
                "검증 상태": "검증됨",
            }
    return None


def get_additional_info_fields(coverage_name):
    field_map = {
        "대인배상Ⅰ": [
            "실제 치료비 또는 손해액",
            "상대방 보험 가입 여부",
        ],
        "대인배상Ⅱ": [
            "실제 치료비 또는 손해액",
            "상대방 보험 가입 여부",
        ],
        "대물배상": [
            "실제 손해액 또는 비용",
            "상대방 보험 가입 여부",
        ],
        "자기신체사고": [
            "실제 치료비 또는 손해액",
            "부상·후유장해 관련 확인자료",
        ],
        "무보험자동차에의한상해": [
            "실제 치료비 또는 손해액",
            "상대방 보험 가입 여부",
        ],
        "자기차량손해": [
            "실제 수리비 또는 손해액",
            "차량가액 또는 손해액",
        ],
    }

    return field_map.get(coverage_name, [])


def parse_optional_amount(value):
    """
    보험증권의 금액 표현을 원 단위 정수로 변환한다.

    예:
    차량가액 1,543만원 -> 15,430,000
    1사고당 10억원 -> 1,000,000,000
    1인당 2억원 한도 -> 200,000,000

    무한·법정한도처럼 고정 숫자가 아닌 표현은 None으로 둔다.
    """
    if value is None:
        return None

    text = str(value).strip()

    if not text:
        return None

    compact = text.replace(",", "").replace(" ", "")

    if "무한" in compact or "법정한도" in compact:
        return None

    # 억원 단위
    match = re.search(r"(\d+(?:\.\d+)?)억원?", compact)
    if match:
        return int(float(match.group(1)) * 100_000_000)

    # 천만원 단위
    match = re.search(r"(\d+(?:\.\d+)?)천만원?", compact)
    if match:
        return int(float(match.group(1)) * 10_000_000)

    # 만원 단위
    match = re.search(r"(\d+(?:\.\d+)?)만원?", compact)
    if match:
        return int(float(match.group(1)) * 10_000)

    # 원 단위
    match = re.search(r"(\d+)원", compact)
    if match:
        return int(match.group(1))

    # 숫자만 들어온 경우에는 기존 입력값과의 호환을 위해 원 단위로 처리
    if re.fullmatch(r"\d+", compact):
        return int(compact)

    return None


def get_coverage_cost_fields(coverage_name):
    field_map = {
        "대인배상Ⅰ": [
            "치료비",
            "수술비",
            "입원비",
            "통원치료비",
            "휴업손해",
        ],
        "대인배상Ⅱ": [
            "치료비",
            "수술비",
            "입원비",
            "통원치료비",
            "휴업손해",
        ],
        "대물배상": [
            "차량수리비",
            "견인비",
            "대차료 또는 교통비",
            "기타 재산손해",
        ],
        "자기신체사고": [
            "치료비",
            "수술비",
            "입원비",
            "통원치료비",
            "휴업손해",
        ],
        "자동차상해": [
            "치료비",
            "수술비",
            "입원비",
            "통원치료비",
            "휴업손해",
        ],
        "무보험자동차에의한상해": [
            "치료비",
            "수술비",
            "입원비",
            "통원치료비",
            "휴업손해",
        ],
        "자기차량손해": [
            "차량수리비",
            "견인비",
        ],
    }
    return field_map.get(coverage_name, [])


def get_cost_input_group(coverage_name):
    if coverage_name in {
        "대인배상Ⅰ",
        "대인배상Ⅱ",
        "자기신체사고",
        "자동차상해",
        "무보험자동차에의한상해",
    }:
        return "사람피해"
    if coverage_name == "대물배상":
        return "상대재산피해"
    if coverage_name == "자기차량손해":
        return "내차량피해"
    return coverage_name


def make_widget_key(*parts):
    normalized = []
    for part in parts:
        text = str(part)
        cleaned = re.sub(
            r"[^0-9A-Za-z가-힣]+",
            "_",
            text,
        ).strip("_")
        if cleaned:
            normalized.append(cleaned.lower())
        else:
            normalized.append("item")
    return "_".join(normalized)


def make_unique_key(screen_area, coverage_code=None, item_code=None, repeat_index=0, question_index=0):
    parts = [str(screen_area).strip() or "screen"]
    if coverage_code:
        parts.append(str(coverage_code).strip())
    if item_code:
        parts.append(str(item_code).strip())
    parts.append(f"r{int(repeat_index or 0)}")
    parts.append(f"q{int(question_index or 0)}")
    return "_".join(parts)


def get_coverage_code(coverage_name):
    return COVERAGE_CODES.get(coverage_name, make_widget_key(coverage_name))


def get_field_code(field_name):
    return COST_FIELD_CODES.get(field_name, make_widget_key(field_name))


def get_coverage_input_key(coverage_name, field_name, index):
    return make_unique_key(
        "actual_cost",
        make_widget_key(get_cost_input_group(coverage_name)),
        get_field_code(field_name),
        repeat_index=index,
        question_index=0,
    )


def get_coverage_selection_key(coverage_name, index):
    return make_unique_key(
        "coverage_select",
        get_coverage_code(coverage_name),
        "selected",
        repeat_index=index,
        question_index=0,
    )


def get_coverage_amount_key(coverage_name, index):
    return make_unique_key(
        "coverage_input",
        get_coverage_code(coverage_name),
        "amount",
        repeat_index=index,
        question_index=0,
    )


def get_coverage_deductible_key(coverage_name, index):
    return make_unique_key(
        "coverage_input",
        get_coverage_code(coverage_name),
        "deductible",
        repeat_index=index,
        question_index=0,
    )


def read_cost_value(coverage_name, field_name, index=0):
    session_inputs = st.session_state.setdefault("actual_cost_inputs", {})
    return session_inputs.get(
        get_coverage_input_key(coverage_name, field_name, index),
        None,
    )


def sum_relevant_costs_for_coverage(coverage_name):
    fields = get_coverage_cost_fields(coverage_name)
    if not fields:
        return None

    total = 0
    for idx, field_name in enumerate(fields):
        value = parse_optional_amount(
            read_cost_value(coverage_name, field_name, idx)
        )
        if value is not None:
            total += value

    return total if total > 0 else None


def get_cost_breakdown_text(coverage_name):
    parts = []
    for index, field_name in enumerate(
        get_coverage_cost_fields(coverage_name)
    ):
        value = parse_optional_amount(
            read_cost_value(coverage_name, field_name, index)
        )
        if value is not None and value > 0:
            parts.append(
                f"{field_name} {format_currency_value(value)}"
            )
    return " + ".join(parts) if parts else "입력 손해액 없음"


def get_injury_grade_limit(injury_grade):
    limits = {
        "1급": 30_000_000,
        "2급": 15_000_000,
        "3급": 12_000_000,
        "4급": 10_000_000,
        "5급": 9_000_000,
        "6급": 7_000_000,
        "7급": 5_000_000,
        "8급": 3_000_000,
        "9급": 2_400_000,
        "10급": 2_000_000,
        "11급": 1_600_000,
        "12급": 1_200_000,
        "13급": 800_000,
        "14급": 500_000,
    }
    return limits.get(str(injury_grade or "").strip())


def sum_treatment_related_costs(coverage_name):
    """치료관계비 최저보장 검토에 사용할 입력값만 합산한다."""
    treatment_fields = {
        "치료비",
        "수술비",
        "입원비",
        "통원치료비",
    }
    total = 0
    has_value = False
    for index, field_name in enumerate(
        get_coverage_cost_fields(coverage_name)
    ):
        if field_name not in treatment_fields:
            continue
        value = parse_optional_amount(
            read_cost_value(coverage_name, field_name, index)
        )
        if value is not None:
            total += value
            has_value = True
    return total if has_value else None


SELF_INJURY_GRADE_LIMITS = {
    15_000_000: {
        "1급": 15_000_000, "2급": 8_000_000, "3급": 7_500_000,
        "4급": 7_000_000, "5급": 5_000_000, "6급": 4_000_000,
        "7급": 2_500_000, "8급": 1_800_000, "9급": 1_400_000,
        "10급": 1_200_000, "11급": 1_200_000, "12급": 1_200_000,
        "13급": 800_000, "14급": 500_000,
    },
    30_000_000: {
        "1급": 30_000_000, "2급": 16_000_000, "3급": 15_000_000,
        "4급": 14_000_000, "5급": 10_000_000, "6급": 8_000_000,
        "7급": 5_000_000, "8급": 3_600_000, "9급": 2_800_000,
        "10급": 2_400_000, "11급": 2_000_000, "12급": 1_800_000,
        "13급": 1_300_000, "14급": 800_000,
    },
    50_000_000: {
        "1급": 50_000_000, "2급": 27_000_000, "3급": 25_000_000,
        "4급": 23_000_000, "5급": 16_500_000, "6급": 13_000_000,
        "7급": 8_000_000, "8급": 6_000_000, "9급": 4_500_000,
        "10급": 4_000_000, "11급": 3_000_000, "12급": 2_900_000,
        "13급": 2_100_000, "14급": 1_300_000,
    },
}


def get_self_injury_grade_limit(injury_grade, limit_text):
    """약관 별표3의 부상 가입금액과 급수별 한도를 찾는다."""
    text = str(limit_text or "")
    injury_segment = ""
    injury_match = re.search(
        r"부상.{0,30}",
        text,
    )
    if injury_match:
        injury_segment = injury_match.group(0)

    injury_plan_limit = parse_optional_amount(injury_segment)
    if injury_plan_limit not in SELF_INJURY_GRADE_LIMITS:
        return None
    return SELF_INJURY_GRADE_LIMITS[injury_plan_limit].get(
        str(injury_grade or "").strip()
    )


def calculate_vehicle_deductible(deductible_text, vehicle_loss):
    """증권상 정률·최저·최고 자기부담금을 현재 차량손해액에 적용한다."""
    text = str(deductible_text or "").strip()
    if not text:
        return 0, False

    percent_match = re.search(r"(\d+(?:\.\d+)?)\s*%", text)
    if percent_match:
        deductible = int(
            max(vehicle_loss or 0, 0)
            * float(percent_match.group(1))
            / 100
        )
        minimum_match = re.search(
            r"최저\s*([0-9.,]+\s*(?:억|천만|만)?원)",
            text,
        )
        maximum_match = re.search(
            r"최고\s*([0-9.,]+\s*(?:억|천만|만)?원)",
            text,
        )
        if minimum_match:
            minimum = parse_optional_amount(minimum_match.group(1))
            if minimum is not None:
                deductible = max(deductible, minimum)
        if maximum_match:
            maximum = parse_optional_amount(maximum_match.group(1))
            if maximum is not None:
                deductible = min(deductible, maximum)
        return deductible, True

    fixed = parse_optional_amount(text)
    return (fixed or 0), fixed is not None


def render_actual_cost_inputs(selected_coverages):
    if not selected_coverages:
        return

    st.subheader("손해액 입력")
    st.caption(
        "같은 사람 피해 금액은 한 번만 입력하면 관련 담보 계산에 함께 사용합니다. "
        "각 칸에는 서로 겹치지 않는 금액을 입력하세요. 치료비가 전체 합계라면 "
        "수술비·입원비·통원치료비 칸은 0으로 두세요."
    )

    rendered_groups = set()
    for item in selected_coverages:
        coverage_name = item.get("담보명")
        group_name = get_cost_input_group(coverage_name)
        if group_name in rendered_groups:
            continue
        fields = get_coverage_cost_fields(coverage_name)
        if not fields:
            continue
        rendered_groups.add(group_name)

        group_labels = {
            "사람피해": "사람 피해 금액",
            "상대재산피해": "상대방 차량·재산 피해 금액",
            "내차량피해": "내 차량 피해 금액",
        }

        heading = group_labels.get(
            group_name,
            f"{coverage_name} 손해 금액",
        )
        st.markdown(
            f'<div class="kh-cost-heading">{escape(heading)}</div>',
            unsafe_allow_html=True,
        )
        for index, field_name in enumerate(fields):
            key = get_coverage_input_key(
                coverage_name,
                field_name,
                index,
            )

            value = render_money_input(
                field_name,
                key,
                help_text="숫자만 입력하세요. 쉼표는 자동으로 표시됩니다.",
            )

            st.session_state.setdefault("actual_cost_inputs", {})[key] = value


def get_total_dedupe_group(coverage_name):
    if coverage_name in {"자기신체사고", "자동차상해"}:
        return "본인 상해담보"
    return coverage_name


def parse_fault_percentage(fault_rate):
    fault_text = str(fault_rate or "")

    match = re.search(
        r"(?<!\d)(\d{1,3})\s*%",
        fault_text,
    )
    if not match:
        return None

    percentage = int(match.group(1))
    return percentage if 0 <= percentage <= 100 else None


def render_input_review_tables(result):
    st.subheader("자동 입력된 가입내용")
    coverage_rows = []
    for item in result.get("selected_coverages", []):
        coverage_rows.append(
            {
                "가입담보명": item.get("담보명", "확인 필요"),
                "가입한도": item.get("가입금액·보상한도", "미입력"),
                "자기부담금": item.get("자기부담금", "미입력"),
                "특약·조건": item.get("특약/비고", "") or "없음",
            }
        )

    if coverage_rows:
        render_white_table(
            ["가입담보명", "가입한도", "자기부담금", "특약·조건"],
            coverage_rows,
        )
    else:
        st.info("선택한 가입담보가 없습니다.")

    with st.expander("입력한 사고정보 확인", expanded=False):
        review_rows = [
            {"항목": "사용자 입장", "입력값": result.get("position") or "미입력"},
            {"항목": "보험 종류", "입력값": result.get("insurance_type") or "미입력"},
            {"항목": "보험회사", "입력값": result.get("insurer") or "미입력"},
            {"항목": "보험기간", "입력값": result.get("insurance_period") or "미입력"},
            {"항목": "부상급수", "입력값": result.get("injury_grade") or "미입력"},
            {"항목": "후유장애급수", "입력값": result.get("disability_grade") or "미입력"},
            {"항목": "입원일수", "입력값": f"{result.get('hospital_days', 0)}일"},
            {"항목": "과실비율", "입력값": result.get("fault_rate") or "미입력"},
            {"항목": "사고내용", "입력값": result.get("accident_text") or "미입력"},
        ]
        render_white_table(["항목", "입력값"], review_rows)


def get_required_policy_inputs(coverage_name):
    base = ["실제 손해비용", "가입한도", "자기부담금", "과실률"]
    if coverage_name in {"대인배상Ⅰ", "대인배상Ⅱ", "자기신체사고", "무보험자동차에의한상해"}:
        base.append("부상급수")
    if coverage_name == "자기신체사고":
        base.append("입원일수")
    return base


def build_policy_analysis_rows(result):
    selected_coverages = result.get("selected_coverages", [])
    if not selected_coverages:
        return [], []

    insurance_type = result.get("insurance_type")
    user_position = result.get("position") or "피해자"
    fault_rate_text = str(result.get("fault_rate") or "")
    parsed_fault = parse_fault_percentage(fault_rate_text)
    applied_fault = parsed_fault if parsed_fault is not None else 0
    provisional_fault = (
        parsed_fault is None
        or any(
            marker in fault_rate_text
            for marker in ["임의", "미확정", "아직 확인되지 않음", "아직 정해지지 않음"]
        )
    )
    injury_grade = result.get("injury_grade") or "아직 확인되지 않음"
    disability_grade = result.get("disability_grade") or "아직 확인되지 않음"
    hospital_days = result.get("hospital_days")
    preconditions = result.get("preconditions") or {}
    injury_grade_limit = get_injury_grade_limit(injury_grade)
    rows = []
    excluded_rows = []
    d1_estimate = 0

    liability_claimant_fault = applied_fault
    if (
        parsed_fault is not None
        and user_position == "사고를 낸 운전자 또는 피보험자"
        and preconditions.get("상대 차량 보험가입 상태") != "상대 차량 없음"
    ):
        liability_claimant_fault = max(0, 100 - applied_fault)

    def apply_limit(amount, limit_value):
        amount = max(int(amount or 0), 0)
        return min(amount, limit_value) if limit_value is not None else amount

    d1_item = next(
        (
            item for item in selected_coverages
            if item.get("담보명") == "대인배상Ⅰ"
        ),
        None,
    )
    shared_injury_loss = sum_relevant_costs_for_coverage("대인배상Ⅰ")
    shared_treatment_loss = sum_treatment_related_costs("대인배상Ⅰ")
    if d1_item and shared_injury_loss is not None:
        d1_limit = parse_optional_amount(
            d1_item.get("가입금액·보상한도")
        )
        available_d1_limits = [
            value for value in [d1_limit, injury_grade_limit]
            if value is not None
        ]
        d1_limit = min(available_d1_limits) if available_d1_limits else None
        d1_adjusted = (
            shared_injury_loss
            * (100 - liability_claimant_fault)
            // 100
        )
        d1_base = max(
            d1_adjusted,
            shared_treatment_loss or 0,
        )
        d1_estimate = apply_limit(d1_base, d1_limit)

    for item in selected_coverages:
        coverage_name = item.get("담보명") or "담보명 확인 필요"
        rule = get_policy_rule(insurance_type, coverage_name)
        limit_text = item.get("가입금액·보상한도", "미입력")
        limit_value = parse_optional_amount(limit_text)
        deductible_value = parse_optional_amount(item.get("자기부담금")) or 0
        actual_total = sum_relevant_costs_for_coverage(coverage_name)
        treatment_total = sum_treatment_related_costs(coverage_name)
        cost_breakdown = get_cost_breakdown_text(coverage_name)
        row_fault_rate = (
            liability_claimant_fault
            if coverage_name in {"대인배상Ⅰ", "대인배상Ⅱ", "대물배상"}
            else applied_fault
        )
        adjusted_loss = (
            actual_total * (100 - row_fault_rate) // 100
            if actual_total is not None
            else 0
        )
        estimated = None
        review_amount = None
        detail = ""
        required = []

        if coverage_name == "대인배상Ⅰ":
            if actual_total is not None:
                available_limits = [
                    value for value in [limit_value, injury_grade_limit]
                    if value is not None
                ]
                effective_limit = min(available_limits) if available_limits else None
                d1_calculation_base = max(
                    adjusted_loss,
                    treatment_total or 0,
                )
                estimated = apply_limit(
                    d1_calculation_base,
                    effective_limit,
                )
                review_amount = estimated
                d1_estimate = estimated
                detail = (
                    f"{cost_breakdown} = 입력 손해액 "
                    f"{format_currency_value(actual_total)} × "
                    f"(100%-피해자측 과실 {row_fault_rate}%) = "
                    f"{format_currency_value(adjusted_loss)}. "
                    f"약관상 치료관계비 최저보장 검토액 "
                    f"{format_currency_value(treatment_total or 0)}과 비교 후 "
                    f"부상급수 {injury_grade}"
                    + (
                        f" 법정 한도 {format_currency_value(injury_grade_limit)} 적용. "
                        if injury_grade_limit is not None
                        else " 법정 한도 확인 필요. "
                    )
                    + f"후유장해급수 {disability_grade}, 입원 {hospital_days or 0}일 입력 반영"
                )
                if effective_limit is None:
                    required.append("부상급수별 대인배상Ⅰ 법정 한도")
            else:
                required.append("치료비·휴업손해 등 실제 손해액")
                detail = "실제 손해액 입력 후 바로 계산"

        elif coverage_name == "대인배상Ⅱ":
            if actual_total is not None:
                fault_adjusted_excess = max(
                    adjusted_loss - d1_estimate,
                    0,
                )
                treatment_expense_excess = max(
                    (treatment_total or 0) - d1_estimate,
                    0,
                )
                remaining = max(
                    fault_adjusted_excess,
                    treatment_expense_excess,
                )
                estimated = apply_limit(remaining, limit_value)
                review_amount = estimated
                detail = (
                    f"{cost_breakdown} = 입력 손해액 "
                    f"{format_currency_value(actual_total)} × "
                    f"(100%-피해자측 과실 {row_fault_rate}%) = "
                    f"{format_currency_value(adjusted_loss)}. "
                    f"여기서 대인배상Ⅰ 예상액 "
                    f"{format_currency_value(d1_estimate)}을 뺀 초과손해와 "
                    f"치료관계비 초과액 {format_currency_value(treatment_expense_excess)}을 "
                    f"비교해 큰 금액을 현재 예상액으로 반영. "
                    f"부상급수 {injury_grade}, 후유장해급수 {disability_grade}, "
                    f"입원 {hospital_days or 0}일 입력 반영"
                )
                if d1_estimate == 0:
                    required.append("대인배상Ⅰ 지급액 또는 법정한도")
            else:
                required.append("치료비·휴업손해 등 실제 손해액")
                required.append("대인배상Ⅰ 지급액")
                detail = "대인배상Ⅰ 초과손해 확인 후 계산"

        elif coverage_name == "대물배상":
            if actual_total is not None:
                estimated = apply_limit(adjusted_loss, limit_value)
                review_amount = estimated
                detail = (
                    f"입력 재산손해 {format_currency_value(actual_total)} × "
                    f"(100%-피해자측 과실 {row_fault_rate}%), 가입한도 범위"
                )
                required.append("교체부품 감가상각 및 약관상 인정기간 확인")
            else:
                required.append("수리비·견인비·대차료 등 실제 손해액")
                detail = "재산손해액 입력 후 바로 계산"

        elif coverage_name == "자기차량손해":
            if actual_total is not None:
                vehicle_deductible, deductible_confirmed = (
                    calculate_vehicle_deductible(
                        item.get("자기부담금"),
                        actual_total,
                    )
                )
                estimated = apply_limit(
                    max(actual_total - vehicle_deductible, 0),
                    limit_value,
                )
                review_amount = estimated
                detail = (
                    f"약관상 피보험자동차 손해액 {format_currency_value(actual_total)} - "
                    f"증권상 자기부담금 {format_currency_value(vehicle_deductible)}. "
                    "자기차량손해에는 사용자 과실비율을 다시 곱하지 않음"
                )
                if not deductible_confirmed:
                    required.append("증권상 자기부담금")
                required.append("잔존물가액·주요 부분품 감가상각 발생 여부")
            else:
                required.append("수리견적서 또는 확정 수리비")
                detail = "차량손해액과 자기부담금 입력 후 계산"

        elif coverage_name == "자기신체사고":
            if actual_total is not None:
                self_injury_limit = get_self_injury_grade_limit(
                    injury_grade,
                    limit_text,
                )
                current_treatment_amount = treatment_total or 0
                estimated = apply_limit(
                    current_treatment_amount,
                    self_injury_limit or limit_value,
                )
                review_amount = estimated
                detail = (
                    f"현재 확인된 치료관계비 {format_currency_value(current_treatment_amount)}. "
                    "약관상 자기신체사고 공제액은 증권의 자기부담금과 다른 개념이므로 "
                    "증권상 자기부담금을 임의로 차감하지 않음. "
                    f"부상급수 {injury_grade}, 후유장해급수 {disability_grade}, "
                    f"입원 {hospital_days or 0}일 입력 반영"
                )
                if self_injury_limit is None:
                    required.append("별표3 부상 가입금액과 급수별 한도 확인")
                required.append("대인배상·무보험차상해·제3자 보상금 등 약관상 공제액")
            else:
                required.append("치료비·휴업손해 등 실제 손해액")
                detail = "실제 손해액 입력 후 바로 계산"
            if injury_grade in {"아직 확인되지 않음", "해당 없음"}:
                required.append("부상급수 확인자료")

        elif coverage_name == "자동차상해":
            if actual_total is not None:
                estimated = apply_limit(actual_total, limit_value)
                review_amount = estimated
                detail = (
                    f"{cost_breakdown} = 현재 확인된 실제손해액 "
                    f"{format_currency_value(actual_total)}. "
                    "자동차상해 약관의 공제액은 증권상 자기부담금과 다른 개념이므로 "
                    "증권상 자기부담금을 임의로 차감하지 않음. "
                    f"부상급수 {injury_grade}, 후유장해급수 {disability_grade}, "
                    f"입원 {hospital_days or 0}일 입력 반영"
                )
                required.append("대인배상·무보험차상해·제3자 보상금 등 약관상 공제액")
            else:
                required.append("치료비·휴업손해 등 실제 손해액")
                detail = "실제 손해액 입력 후 바로 계산"
            if injury_grade in {"아직 확인되지 않음", "해당 없음"}:
                required.append("부상급수 확인자료")

        elif coverage_name == "무보험자동차에의한상해":
            policy_cost_label = (
                "약관상 비용(손해방지·경감비용 및 권리보전·행사비용)"
            )
            deduction_labels = [
                "공제 1. 대인배상Ⅰ·책임공제·정부보장사업에서 지급될 수 있는 금액",
                "공제 2. 배상의무자 차량의 대인배상Ⅱ·공제계약에서 지급될 수 있는 금액",
                "공제 3. 탑승 차량의 대인배상Ⅱ·공제계약에서 지급될 수 있는 금액",
                "공제 4. 배상의무자에게 이미 받은 손해배상금",
                "공제 5. 제3자가 부담할 금액 중 이미 받은 금액",
            ]
            policy_cost = parse_optional_amount(
                preconditions.get(policy_cost_label)
            )
            deduction_values = [
                parse_optional_amount(preconditions.get(label))
                for label in deduction_labels
            ]
            verified_policy_formula = (
                rule is not None
                and rule.get("검증 상태") == "검증됨"
                and "보험금지급기준" in str(rule.get("계산식") or "")
                and "공제액" in str(rule.get("계산식") or "")
            )
            required_coverages = {
                "대인배상Ⅰ",
                "대인배상Ⅱ",
                "대물배상",
                "자기신체사고",
            }
            joined_coverages = {
                selected.get("담보명")
                for selected in selected_coverages
            }

            if not verified_policy_formula:
                required.append("적용 보험종류의 약관 PDF 계산식 검증")
            if actual_total is None:
                required.append("치료비·휴업손해 등 약관상 손해액")
            if preconditions.get("상대 차량 보험가입 상태") != "무보험":
                required.append("상대 차량이 무보험자동차라는 확인자료")
            if preconditions.get("무보험자동차상해 적용 검토") != "적용 검토 필요":
                required.append("무보험자동차상해 적용 대상 확인")
            if not required_coverages.issubset(joined_coverages):
                required.append(
                    "대인배상Ⅰ·Ⅱ, 대물배상, 자기신체사고 가입 확인"
                )
            if policy_cost is None:
                required.append("약관상 비용 금액 확인")
            for label, value in zip(deduction_labels, deduction_values):
                if value is None:
                    required.append(label + " 확인")
            if provisional_fault:
                required.append("최종 과실비율")

            if verified_policy_formula and actual_total is not None:
                reflected_policy_cost = policy_cost or 0
                reflected_deductions = sum(
                    value for value in deduction_values
                    if value is not None
                )
                policy_formula_amount = max(
                    adjusted_loss
                    + reflected_policy_cost
                    - reflected_deductions,
                    0,
                )
                estimated = apply_limit(
                    policy_formula_amount,
                    limit_value,
                )
                review_amount = estimated
                detail = (
                    "약관 지급보험금 계산식 적용: "
                    f"{cost_breakdown} = 입력 사람피해액 "
                    f"{format_currency_value(actual_total)} × "
                    f"(100%-피보험자 과실 {row_fault_rate}%) = "
                    f"{format_currency_value(adjusted_loss)} "
                    f"+ 현재 확인된 약관상 비용 {format_currency_value(reflected_policy_cost)} "
                    f"- 현재 확인된 공제액 합계 {format_currency_value(reflected_deductions)} "
                    f"= {format_currency_value(policy_formula_amount)}. "
                    "보험증권 가입한도 적용. 빈칸인 비용·공제항목은 0원으로 "
                    "확정한 것이 아니라 아직 반영하지 않은 변수이므로, 확인 후 "
                    "입력하면 예상액이 다시 계산됩니다."
                )
            else:
                detail = (
                    "현재 입력한 사람피해액이 없거나 적용 보험종류의 약관 계산식이 "
                    "확인되지 않아 이 담보만 계산하지 않았습니다. 손해액 또는 약관 "
                    "근거가 확인되면 즉시 예상액에 포함됩니다."
                )

        else:
            required.append("해당 특약의 지급요건과 손해자료")
            detail = "가입내용은 확인됨 · 금액형 담보인지 추가 확인"

        if provisional_fault and coverage_name in {
            "대인배상Ⅰ", "대인배상Ⅱ", "대물배상",
            "무보험자동차에의한상해", "자기차량손해",
        }:
            required.append("최종 과실비율")

        if coverage_name in {
            "대인배상Ⅰ", "대인배상Ⅱ", "자기신체사고",
            "자동차상해", "무보험자동차에의한상해",
        }:
            if disability_grade in {"아직 확인되지 않음", "해당 없음"}:
                required.append("후유장해가 남는 경우 장해진단서")
            if hospital_days in {None, 0}:
                required.append("입원한 경우 입·퇴원확인서")

        required = list(dict.fromkeys(required))
        amount_text = (
            format_currency_value(review_amount)
            if review_amount is not None
            else "계산 보류"
        )
        note_parts = [
            detail,
            "입력 및 추정 데이터를 바탕으로 산정한 예상 금액이며 최종 확정 금액이 아닙니다.",
        ]
        if required:
            note_parts.append("추가 확인: " + ", ".join(required))

        row_values = {
            "가입담보명": coverage_name,
            "가입한도": limit_text or "미입력",
            "청구 검토금액": amount_text,
            "비고 및 세부 산정 내역": " | ".join(note_parts),
            "실제 손해비용": (
                format_currency_value(actual_total)
                if actual_total is not None
                else "미입력"
            ),
            "자기부담금": item.get("자기부담금", "미입력") or "미입력",
            "후유장애급수": disability_grade,
            "부상급수": injury_grade,
            "입원일수": f"{hospital_days or 0}일",
            "적용 과실비율": (
                f"{row_fault_rate}%"
                + (" 가정" if provisional_fault else "")
            ),
            "약관상 예상 지급액": amount_text,
            "적용 약관 규정 및 산정근거": detail,
            "비고": "추가 확인: " + ", ".join(required) if required else "현재 입력값으로 1차 계산",
            "계산가능금액": estimated,
            "검토가능금액": review_amount,
            "추가 확인": required,
            "근거 PDF": (
                "개인용 약관.pdf"
                if coverage_name == "무보험자동차에의한상해"
                and insurance_type == "개인용 자동차보험"
                else rule.get("PDF 파일명", "확인 필요") if rule else "확인 필요"
            ),
            "PDF 페이지": (
                "52-53"
                if coverage_name == "무보험자동차에의한상해"
                and insurance_type == "개인용 자동차보험"
                else rule.get("PDF 페이지", "확인 필요") if rule else "확인 필요"
            ),
            "조문명": (
                "제20조(지급보험금의 계산)"
                if coverage_name == "무보험자동차에의한상해"
                and insurance_type == "개인용 자동차보험"
                else rule.get("조문명", "확인 필요") if rule else "확인 필요"
            ),
        }
        rows.append(row_values)

        if review_amount is None:
            excluded_rows.append({
                "담보명": coverage_name,
                "사유": ", ".join(required) or "손해액 입력 필요",
                "근거 PDF": row_values["근거 PDF"],
                "PDF 페이지": row_values["PDF 페이지"],
                "조문명": row_values["조문명"],
            })

    total_estimate = sum(
        row.get("검토가능금액") or 0
        for row in rows
    )
    st.session_state["policy_analysis_payload"] = {
        "result": result,
        "rows": rows,
        "excluded_rows": excluded_rows,
        "total_estimate": total_estimate,
    }
    return rows, excluded_rows


def build_analysis_table(rows):
    visible_headers = [
        "담보 구분",
        "가입 한도",
        "청구 검토 금액 (예상)",
        "비고 및 세부 산정 내역",
    ]
    table_rows = []
    amount_by_group = {}

    for row in rows:
        amount = row.get("검토가능금액")
        group = get_total_dedupe_group(row.get("가입담보명", ""))
        amount_by_group[group] = max(
            amount_by_group.get(group, 0),
            amount or 0,
        )
        table_rows.append({
            "담보 구분": row.get("가입담보명", ""),
            "가입 한도": row.get("가입한도", ""),
            "청구 검토 금액 (예상)": row.get("청구 검토금액", "0원"),
            "비고 및 세부 산정 내역": row.get("비고 및 세부 산정 내역", ""),
        })

    total_estimate = sum(amount_by_group.values())
    table_rows.append({
        "담보 구분": "현재 입력값 기준 예상 보험금 합계",
        "가입 한도": "",
        "청구 검토 금액 (예상)": format_currency_value(total_estimate),
        "비고 및 세부 산정 내역": (
            "현재 입력값으로 계산한 금액의 합계입니다. 선택 관계인 자기신체사고와 "
            "자동차상해는 큰 금액 하나만 반영했습니다. "
            "담보 간 중복보상 제한, 이미 받은 보험금, 면책·공제와 최종 과실비율이 "
            "확인되면 금액이 달라질 수 있습니다."
        ),
    })
    return table_rows, total_estimate


def build_analysis_excel(result, rows, excluded_rows, total_estimate):
    try:
        from io import BytesIO
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Side, Border, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl이 설치되어 있지 않습니다.") from exc

    workbook = Workbook()

    input_sheet = workbook.active
    input_sheet.title = "입력정보"
    input_sheet.append(["항목", "값"])
    input_sheet.append(["사용자 입장", result.get("position", "미입력")])
    input_sheet.append(["보험 종류", result.get("insurance_type", "미입력")])
    input_sheet.append(["보험회사", result.get("insurer", "미입력")])
    input_sheet.append(["보험기간", result.get("insurance_period", "미입력")])
    input_sheet.append(["부상급수", result.get("injury_grade", "미입력")])
    input_sheet.append(["후유장애급수", result.get("disability_grade", "미입력")])
    input_sheet.append(["입원일수", result.get("hospital_days", 0)])
    input_sheet.append(["적용 과실비율", result.get("fault_rate", "미입력")])
    input_sheet.append(["사고내용", result.get("accident_text", "미입력")])

    for cell in input_sheet["A1:B1"][0]:
        cell.font = Font(bold=True)

    input_sheet.column_dimensions["A"].width = 22
    input_sheet.column_dimensions["B"].width = 55

    policy_rows, _ = build_analysis_table(rows)
    policy_sheet = workbook.create_sheet("보상분석표")

    policy_columns = [
        ("담보 구분", "담보 구분"),
        ("가입 한도", "가입 한도"),
        ("청구 검토 금액 (예상)", "청구 검토 금액 (예상)"),
        ("비고 및 세부 산정 내역", "비고 및 세부 산정 내역"),
    ]

    policy_sheet.append([label for label, key in policy_columns])

    for row in policy_rows:
        policy_sheet.append([
            row.get(key, "")
            for label, key in policy_columns
        ])

    title_fill = PatternFill("solid", fgColor="D9E2F3")
    total_fill = PatternFill("solid", fgColor="F3F6F9")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in policy_sheet[1]:
        cell.font = Font(bold=True, color="000000")
        cell.fill = title_fill
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row_cells in policy_sheet.iter_rows(
        min_row=2,
        max_row=policy_sheet.max_row,
        min_col=1,
        max_col=len(policy_columns),
    ):
        for cell in row_cells:
            cell.border = border
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    if policy_sheet.max_row > 1:
        last_row = policy_sheet[policy_sheet.max_row]
        for cell in last_row:
            cell.font = Font(bold=True)
            cell.fill = total_fill

    widths = [30, 22, 25, 95]
    for index, width in enumerate(widths, start=1):
        policy_sheet.column_dimensions[get_column_letter(index)].width = width

    policy_sheet.freeze_panes = "A2"
    policy_sheet.print_title_rows = "1:1"
    policy_sheet.page_setup.orientation = "landscape"
    policy_sheet.page_setup.fitToWidth = 1
    policy_sheet.page_setup.fitToHeight = 0
    policy_sheet.sheet_view.showGridLines = True

    consultation_sheet = workbook.create_sheet("상담내용")
    consultation_sheet.append([
        "번호",
        "질문",
        "관련 담보",
        "약관 답변",
        "조문",
        "PDF",
        "페이지",
    ])

    evidence_sheet = workbook.create_sheet("약관근거")
    evidence_sheet.append([
        "담보명",
        "근거 PDF",
        "PDF 페이지",
        "조문명",
    ])
    for row in rows:
        evidence_sheet.append([
            row.get("가입담보명", ""),
            row.get("근거 PDF", "확인 필요"),
            row.get("PDF 페이지", "확인 필요"),
            row.get("조문명", "확인 필요"),
        ])

    excluded_sheet = workbook.create_sheet("계산제외항목")
    excluded_sheet.append([
        "담보명",
        "사유",
        "근거 PDF",
        "PDF 페이지",
        "조문명",
    ])
    for row in excluded_rows:
        excluded_sheet.append([
            row.get("담보명", ""),
            row.get("사유", ""),
            row.get("근거 PDF", ""),
            row.get("PDF 페이지", ""),
            row.get("조문명", ""),
        ])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_policy_analysis_result(result):
    selected_coverages = result.get("selected_coverages", [])
    if not selected_coverages:
        st.info("가입담보를 먼저 확인해 주세요.")
        return

    rows, excluded_rows = build_policy_analysis_rows(result)
    analysis_table, total_estimate = build_analysis_table(rows)

    st.subheader("입력값 기준 보상금 산정 요약")
    st.caption(
        "보험증권에서 확인한 가입내용과 현재 입력한 손해액을 기준으로 계산했습니다. "
        "각 담보의 추가 확인사항을 반영하면 예상 금액이 다시 계산됩니다."
    )

    render_white_table(
        [
            "담보 구분",
            "가입 한도",
            "청구 검토 금액 (예상)",
            "비고 및 세부 산정 내역",
        ],
        analysis_table,
    )

    st.markdown(
        f"""
        <div class="kh-result-card">
            <div class="kh-result-label">현재 입력값 기준 예상 보험금 합계</div>
            <div class="kh-result-amount">{escape(format_currency_value(total_estimate))}</div>
            <div class="kh-result-note">
                현재 입력값으로 계산한 1차 예상액입니다. 담보 간 중복보상 제한,
                이미 받은 보험금, 최종 과실비율과 보험회사 심사를 반영하면 달라질 수 있습니다.
                미입력 변수는 0원으로 확정한 것이 아니라 현재 예상액에 아직 반영하지 않은 값입니다.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    action_items = []
    fault_rate_text = str(result.get("fault_rate") or "")
    if (
        parse_fault_percentage(fault_rate_text) is None
        or any(
            marker in fault_rate_text
            for marker in ["임의", "미확정", "아직 정해지지 않음", "아직 확인되지 않음"]
        )
    ):
        action_items.append(
            "최종 과실비율: 보험회사 협의서, 과실비율 분쟁심의 결과 또는 판결·조정자료"
        )

    injury_names = {
        "대인배상Ⅰ",
        "대인배상Ⅱ",
        "자기신체사고",
        "자동차상해",
        "무보험자동차에의한상해",
    }
    selected_names = {
        item.get("담보명")
        for item in selected_coverages
    }
    if selected_names & injury_names:
        action_items.extend([
            "치료 관련 자료: 진단서, 진료비 계산서·영수증, 입·퇴원확인서와 향후 통원계획",
            "소득 관련 자료: 급여명세서, 소득금액증명원 또는 과세표준증명원",
        ])
        if result.get("disability_grade") in {
            None, "", "아직 확인되지 않음", "해당 없음"
        }:
            action_items.append(
                "후유장해가 남는 경우: 치료 종결 후 발급된 후유장해진단서와 장해평가 자료"
            )

    if "무보험자동차에의한상해" in selected_names:
        action_items.append(
            "무보험자동차상해 검토: 상대 차량 보험가입사실과 대인배상 가입 범위"
        )

    if "자기차량손해" in selected_names:
        action_items.append(
            "차량 손해 자료: 수리견적서·정비명세서·사고사진과 증권상 자기부담금"
        )

    for row in rows:
        for item in row.get("추가 확인", []):
            if item not in action_items:
                action_items.append(item)

    action_items = list(dict.fromkeys(action_items))
    if action_items:
        items_html = "".join(
            f"<li>{escape(str(item))}</li>"
            for item in action_items
        )
        st.markdown("#### 금액을 더 정확하게 만들기 위해 준비할 자료")
        st.markdown(
            f"""
            <div class="kh-result-card">
                <div class="kh-result-label">다음 행동</div>
                <ol class="kh-action-list">{items_html}</ol>
                <div class="kh-result-note">
                    자료를 확인한 뒤 위 입력값을 고치고 다시 계산하면 예상 금액이 즉시 갱신됩니다.
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with st.expander("계산에 적용한 조건과 약관 근거 보기", expanded=False):
        detail_rows = []
        for row in rows:
            detail_rows.append({
                "담보": row.get("가입담보명", ""),
                "손해액": row.get("실제 손해비용", ""),
                "자기부담금": row.get("자기부담금", ""),
                "부상·장해·입원": (
                    f"{row.get('부상급수', '')} · "
                    f"{row.get('후유장애급수', '')} · "
                    f"{row.get('입원일수', '')}"
                ),
                "과실비율": row.get("적용 과실비율", ""),
                "약관 근거": (
                    f"{row.get('조문명', '확인 필요')} · "
                    f"{row.get('근거 PDF', '확인 필요')} "
                    f"{row.get('PDF 페이지', '확인 필요')}쪽"
                ),
            })
        render_white_table(
            ["담보", "손해액", "자기부담금", "부상·장해·입원", "과실비율", "약관 근거"],
            detail_rows,
        )

    payload = st.session_state.get("policy_analysis_payload")
    if payload:
        try:
            excel_data = build_analysis_excel(
                payload["result"],
                payload["rows"],
                payload["excluded_rows"],
                total_estimate,
            )
            st.download_button(
                label="보상금 산정표 엑셀로 받기",
                data=excel_data,
                file_name="입력값_기준_보상금_산정표.xlsx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                key="download_policy_analysis_excel_v3",
                width="content",
            )
        except RuntimeError as exc:
            st.warning(str(exc))


# =========================================================
# 보험증권 사진 자동 판독
# =========================================================

def normalize_detected_insurance_type(value):
    text = str(value or "").replace(" ", "")
    if "이륜" in text or "오토바이" in text:
        return "이륜차 자동차보험"
    if "영업" in text:
        return "영업용 자동차보험"
    if "업무" in text:
        return "업무용 자동차보험"
    if "개인" in text:
        return "개인용 자동차보험"
    return None


def normalize_ocr_coverage_name(value):
    name = str(value or "").strip()
    mapping = {
        "무보험자동차상해": "무보험자동차에의한상해",
        "무보험차상해": "무보험자동차에의한상해",
        "무보험자동차에의한상해": "무보험자동차에의한상해",
        "자동차상해특약": "자동차상해",
        "자동차상해": "자동차상해",
    }
    return mapping.get(name, name)


def apply_policy_ocr_to_form(validated):
    """
    보험증권 OCR 결과를 화면 입력값에 연결한다.

    원칙:
    - 보험증권 원문을 보존한다.
    - 내부 분석용 표준 담보명은 별도로 사용한다.
    - 숫자로 변환되지 않는 '무한', '법정한도' 등의 원문도 버리지 않는다.
    - OCR에서 확인되지 않은 담보를 임의로 선택하지 않는다.
    """
    basic = validated.get("basic_info") or {}

    detected_type = normalize_detected_insurance_type(
        basic.get("보험종류")
        or basic.get("차량종류")
        or basic.get("보험상품명")
    )

    if detected_type:
        st.session_state["insurance_type_select"] = detected_type
    else:
        detected_type = st.session_state.get(
            "insurance_type_select",
            INSURANCE_TYPE_OPTIONS[0],
        )

    storage_key = make_storage_key(detected_type)

    insurer_key = (
        f"{storage_key}_insurer_v2_20260925_001"
    )
    product_key = (
        f"{storage_key}_product_20260928_001"
    )
    period_key = (
        f"{storage_key}_period_20260925_001"
    )

    if basic.get("보험회사"):
        st.session_state[insurer_key] = str(
            basic.get("보험회사")
        ).strip()

    if basic.get("보험상품명"):
        st.session_state[product_key] = str(
            basic.get("보험상품명")
        ).strip()

    if basic.get("보험기간"):
        st.session_state[period_key] = str(
            basic.get("보험기간")
        ).strip()

    coverage_names = COVERAGES_BY_TYPE.get(
        detected_type,
        [],
    )

    meta = {}
    source_coverages = []

    for item in validated.get("coverages") or []:
        if not isinstance(item, dict):
            continue

        original_name = str(
            item.get("original_coverage_name")
            or ""
        ).strip()

        standard = normalize_ocr_coverage_name(
            item.get("standard_coverage_name")
            or original_name
        )

        original_limit = str(
            item.get("original_limit")
            or ""
        ).strip()

        deductible_original = str(
            item.get("deductible")
            or ""
        ).strip()

        special_conditions = str(
            item.get("special_conditions")
            or ""
        ).strip()

        source_text = str(
            item.get("source_text")
            or ""
        ).strip()

        limit_won = item.get(
            "limit_amount_won"
        )

        deductible_won = parse_amount_to_won(
            deductible_original
        )

        source_record = {
            "original_coverage_name": original_name,
            "standard_coverage_name": standard,
            "original_limit": original_limit,
            "limit_amount_won": limit_won,
            "deductible": deductible_original,
            "deductible_amount_won": deductible_won,
            "special_conditions": special_conditions,
            "source_text": source_text,
            "confidence": item.get("confidence"),
            "needs_review": bool(
                item.get("needs_review")
            ),
        }

        source_coverages.append(
            source_record
        )

        # 현재 앱의 표준 담보와 연결되지 않는 항목도
        # source_coverages에는 그대로 보존한다.
        if standard not in coverage_names:
            continue

        index = coverage_names.index(
            standard
        )

        st.session_state[
            get_coverage_selection_key(
                standard,
                index,
            )
        ] = True

        # 숫자로 명확하게 변환되는 경우에는
        # 기존 금액 입력칸에도 자동 입력한다.
        if (
            isinstance(limit_won, int)
            and limit_won >= 0
        ):
            st.session_state[
                get_coverage_amount_key(
                    standard,
                    index,
                )
            ] = str(limit_won)

        if (
            isinstance(deductible_won, int)
            and deductible_won >= 0
        ):
            st.session_state[
                get_coverage_deductible_key(
                    standard,
                    index,
                )
            ] = str(deductible_won)

        notes = []

        if original_name and original_name != standard:
            notes.append(
                f"증권 표기: {original_name}"
            )

        if (
            original_limit
            and limit_won is None
        ):
            notes.append(
                f"가입금액·보상한도 원문: {original_limit}"
            )

        if (
            deductible_original
            and deductible_won is None
        ):
            notes.append(
                f"자기부담금 원문: {deductible_original}"
            )

        if special_conditions:
            notes.append(
                f"특약·조건: {special_conditions}"
            )

        if item.get("needs_review"):
            notes.append(
                "사진 판독 확인 필요"
            )

        meta[standard] = {
            "original_name": (
                original_name or standard
            ),
            "original_limit": original_limit,
            "limit_amount_won": limit_won,
            "deductible_original": (
                deductible_original
            ),
            "deductible_amount_won": (
                deductible_won
            ),
            "special_conditions": (
                special_conditions
            ),
            "source_text": source_text,
            "notes": " · ".join(notes),
            "needs_review": bool(
                item.get("needs_review")
            ),
        }

    # 다음 단계의 '보험증권 입력내용 확인·수정' 화면에서
    # 증권 원문 전체를 사용할 수 있도록 별도로 보관한다.
    st.session_state[
        "policy_ocr_source_coverages"
    ] = source_coverages

    # 기존 보상분석과의 호환성을 위해 유지한다.
    st.session_state[
        "policy_ocr_coverage_meta"
    ] = meta


def process_policy_image_upload(uploaded_image):
    try:
        image_bytes = uploaded_image.getvalue()
    except Exception as exc:
        st.session_state["policy_ocr_error"] = f"사진 파일을 읽지 못했습니다: {exc}"
        return False

    if not image_bytes:
        st.session_state["policy_ocr_error"] = "보험증권 사진이 비어 있습니다. 다시 촬영해 주세요."
        return False

    image_hash = hashlib.sha256(image_bytes).hexdigest()
    if (
        image_hash == st.session_state.get("policy_ocr_image_hash")
        and st.session_state.get("policy_ocr_result")
    ):
        apply_policy_ocr_to_form(st.session_state["policy_ocr_result"])
        return True

    mime_type = getattr(uploaded_image, "type", "") or "image/jpeg"
    suffix_map = {
        "image/jpeg": ".jpg",
        "image/jpg": ".jpg",
        "image/png": ".png",
        "image/webp": ".webp",
        "image/heic": ".heic",
        "image/heif": ".heif",
    }

    # 브라우저가 HEIC/HEIF MIME 형식을 정확히 보내지 않는 경우에도
    # 원래 파일명의 확장자를 확인하여 원본 형식을 보존한다.
    original_name = str(
        getattr(uploaded_image, "name", "") or ""
    )
    original_suffix = Path(original_name).suffix.lower()

    if original_suffix in {
        ".jpg",
        ".jpeg",
        ".png",
        ".webp",
        ".heic",
        ".heif",
    }:
        suffix = original_suffix
    else:
        suffix = suffix_map.get(
            mime_type.lower(),
            ".jpg",
        )
    temp_path = ""
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
            temp_file.write(image_bytes)
            temp_path = temp_file.name

        raw_result = read_policy_image(
            temp_path,
            api_key=get_openai_key(),
        )
        validated = validate_policy_result(raw_result)
        st.session_state["policy_ocr_result"] = validated
        st.session_state["policy_ocr_image_hash"] = image_hash
        st.session_state["policy_ocr_error"] = ""
        apply_policy_ocr_to_form(validated)
        return True
    except Exception as exc:
        st.session_state["policy_ocr_error"] = str(exc)
        return False
    finally:
        if temp_path:
            try:
                os.unlink(temp_path)
            except OSError:
                pass


def infer_accident_preconditions(accident_text):
    text = str(accident_text or "")
    inferred = {}
    rules = [
        (r"무면허", "사고 운전자 면허 상태", "무면허"),
        (r"면허\s*정지", "사고 운전자 면허 상태", "면허정지 중"),
        (r"면허\s*취소", "사고 운전자 면허 상태", "면허취소 상태"),
        (r"음주\s*운전|술을?\s*(마시|먹).{0,12}운전", "음주운전 여부", "예"),
        (r"마약|약물\s*운전", "마약·약물운전 여부", "예"),
        (r"뺑소니|사고.{0,8}도주|조치.{0,8}없이.{0,8}떠", "사고발생 후 조치의무 위반 여부", "예"),
    ]
    for pattern, label, value in rules:
        if re.search(pattern, text, flags=re.IGNORECASE):
            inferred[label] = value
    return inferred


initialize_state()

# =========================================================
# 홈페이지 시작 화면
# =========================================================

if not st.session_state.get("homepage_started", False):
    st.html(
        """
        <style>
        .kh-home {
            --kh-ink: #172033;
            --kh-sub: #596579;
            --kh-blue: #2563eb;
            --kh-blue-dark: #1746b0;
            --kh-line: #dfe6ef;
            --kh-soft: #f5f8fc;
        }

        .kh-hero {
            position: relative;
            overflow: hidden;
            padding: 4.4rem 3.2rem 3.8rem;
            margin: 0.5rem 0 2.2rem;
            border: 1px solid var(--kh-line);
            border-radius: 28px;
            background:
                radial-gradient(circle at 85% 15%, #dce9ff 0, transparent 33%),
                linear-gradient(145deg, #ffffff 0%, #f5f8fc 100%);
            box-shadow: 0 18px 55px rgba(29, 54, 92, 0.08);
        }

        .kh-brand {
            display: inline-flex;
            align-items: center;
            gap: 0.55rem;
            padding: 0.45rem 0.8rem;
            margin-bottom: 1.5rem;
            border: 1px solid #cfdbed;
            border-radius: 999px;
            background: rgba(255, 255, 255, 0.85);
            color: #355174;
            font-size: 0.88rem;
            font-weight: 700;
        }

        .kh-compass {
            display: inline-grid;
            place-items: center;
            width: 1.7rem;
            height: 1.7rem;
            border-radius: 50%;
            background: var(--kh-blue);
            color: #ffffff;
            font-size: 0.9rem;
        }

        .kh-hero h1 {
            max-width: 760px;
            margin: 0;
            color: var(--kh-ink);
            font-size: clamp(2.35rem, 5.5vw, 4.25rem);
            line-height: 1.16;
            letter-spacing: -0.055em;
        }

        .kh-highlight {
            color: var(--kh-blue);
        }

        .kh-hero p {
            max-width: 650px;
            margin: 1.5rem 0 0;
            color: var(--kh-sub);
            font-size: 1.12rem;
            line-height: 1.85;
            word-break: keep-all;
        }

        .kh-trust-row {
            display: flex;
            flex-wrap: wrap;
            gap: 0.65rem;
            margin-top: 1.6rem;
        }

        .kh-trust {
            padding: 0.48rem 0.72rem;
            border-radius: 9px;
            background: #ffffff;
            border: 1px solid var(--kh-line);
            color: #43516a;
            font-size: 0.86rem;
            font-weight: 650;
        }

        .kh-section-title {
            margin: 2.6rem 0 0.45rem;
            color: var(--kh-ink);
            font-size: 1.65rem;
            font-weight: 800;
            letter-spacing: -0.035em;
        }

        .kh-section-copy {
            margin-bottom: 1.35rem;
            color: var(--kh-sub);
            line-height: 1.7;
        }

        .kh-steps {
            display: grid;
            grid-template-columns: repeat(3, minmax(0, 1fr));
            gap: 1rem;
            margin-bottom: 2rem;
        }

        .kh-step {
            min-height: 175px;
            padding: 1.4rem;
            border: 1px solid var(--kh-line);
            border-radius: 18px;
            background: #ffffff;
        }

        .kh-step-number {
            display: inline-grid;
            place-items: center;
            width: 2.1rem;
            height: 2.1rem;
            margin-bottom: 1rem;
            border-radius: 11px;
            background: #eaf1ff;
            color: var(--kh-blue-dark);
            font-weight: 850;
        }

        .kh-step h3 {
            margin: 0 0 0.55rem;
            color: var(--kh-ink);
            font-size: 1.08rem;
        }

        .kh-step p {
            margin: 0;
            color: var(--kh-sub);
            font-size: 0.94rem;
            line-height: 1.65;
            word-break: keep-all;
        }

        .kh-note {
            margin: 1.2rem 0 0;
            padding: 1rem 1.15rem;
            border-left: 4px solid #9db9ed;
            border-radius: 8px;
            background: var(--kh-soft);
            color: #526076;
            font-size: 0.88rem;
            line-height: 1.65;
        }

        div.stButton > button[kind="secondary"] {
            min-height: 3.5rem;
            border-radius: 14px !important;
            background: var(--kh-blue) !important;
            border-color: var(--kh-blue) !important;
            font-size: 1.05rem;
            box-shadow: 0 10px 26px rgba(37, 99, 235, 0.18);
        }

        div.stButton > button[kind="secondary"]:hover {
            background: var(--kh-blue-dark) !important;
            border-color: var(--kh-blue-dark) !important;
        }


        .kh-click-cue {
            display: inline-flex;
            align-items: center;
            justify-content: center;
            gap: 0.6rem;
            margin-top: 1.15rem;
            padding: 0.7rem 1rem;
            border: 1px solid #bfdbfe;
            border-radius: 999px;
            background: #eff6ff;
            color: #1d4ed8;
            font-size: 0.95rem;
            font-weight: 750;
        }

        .kh-click-cue .kh-pointer {
            font-size: 1.35rem;
        }

        .kh-cta-guide {
            max-width: 680px;
            margin: 2.5rem auto 1.1rem;
            text-align: center;
        }

        .kh-cta-guide h2 {
            margin: 0 0 0.6rem;
            color: #172033;
            font-size: 1.55rem;
            letter-spacing: -0.035em;
        }

        .kh-cta-guide p {
            margin: 0;
            color: #667085;
            font-size: 0.96rem;
            line-height: 1.7;
            word-break: keep-all;
        }

        div[data-testid="stButton"] {
            max-width: 680px;
            margin: 0 auto;
        }

        div[data-testid="stButton"] > button {
            min-height: 3.8rem !important;
            border: 1px solid #2563eb !important;
            border-radius: 16px !important;
            background: #2563eb !important;
            color: #ffffff !important;
            font-size: 1.08rem !important;
            font-weight: 800 !important;
            box-shadow: 0 12px 28px rgba(37, 99, 235, 0.22) !important;
        }

        div[data-testid="stButton"] > button:hover {
            background: #1746b0 !important;
            border-color: #1746b0 !important;
            transform: translateY(-1px);
        }

        div[data-testid="stButton"] > button,
        div[data-testid="stButton"] > button p,
        div[data-testid="stButton"] > button span {
            color: #ffffff !important;
            opacity: 1 !important;
            cursor: pointer !important;
        }

        @media (max-width: 720px) {
            .kh-hero {
                padding: 2.7rem 1.35rem 2.4rem;
                border-radius: 21px;
            }

            .kh-hero h1 {
                font-size: 2.45rem;
            }

            .kh-hero p {
                font-size: 1rem;
            }

            .kh-steps {
                grid-template-columns: 1fr;
            }

            .kh-step {
                min-height: auto;
            }
        }
        </style>

        <div class="kh-home">
            <section class="kh-hero">
                <div class="kh-brand">
                    <span class="kh-compass">↗</span>
                    기칸의 자동차보험 보상 나침반
                </div>

                <h1>
                    사고 뒤,<br>
                    <span class="kh-highlight">약관에 근거한</span><br>
                    예상 보험금
                </h1>

                <p>
                    사고내용을 말하고 보험증권을 촬영하세요.
                    보험회사·보험기간·가입담보를 사진에서 자동으로 읽어 채운 뒤,
                    사용자는 입력된 내용을 확인하고 수정하면 됩니다.
                </p>

                <div class="kh-trust-row">
                    <span class="kh-trust">음성으로 간편 입력</span>
                    <span class="kh-trust">보험증권 사진 자동입력</span>
                    <span class="kh-trust">약관 원문과 페이지 확인</span>
                </div>
            </section>

            <div class="kh-section-title">세 단계로 확인합니다</div>
            <div class="kh-section-copy">
                많은 정보를 한꺼번에 보여주지 않고, 사고와 가입담보에 필요한 내용부터 확인합니다.
            </div>

            <div class="kh-steps">
                <div class="kh-step">
                    <div class="kh-step-number">1</div>
                    <h3>사고내용 말하기</h3>
                    <p>마이크로 설명하면 음성이 문자로 바뀌며, 잘못 변환된 부분은 직접 수정할 수 있습니다.</p>
                </div>

                <div class="kh-step">
                    <div class="kh-step-number">2</div>
                    <h3>보험증권 촬영하기</h3>
                    <p>보험증권을 사진으로 찍으면 보험회사·보험기간·가입담보와 한도를 자동으로 읽어 입력합니다.</p>
                </div>

                <div class="kh-step">
                    <div class="kh-step-number">3</div>
                    <h3>확인하고 분석하기</h3>
                    <p>자동 입력된 내용을 사용자가 확인한 뒤 사고내용과 담보를 연결해 관련 약관과 보상 항목을 검토합니다.</p>
                </div>
            </div>
        </div>
        """
    )

    st.html(
        """
        <div class="kh-cta-guide">
            <h2>내 보험의 보상 내용을 확인해 보세요</h2>
        </div>
        """
    )

    if st.button(
        "🖱️  여기를 눌러 보상 확인 시작  →",
        key="homepage_start_button",
        type="primary",
        width="stretch",
    ):
        st.session_state.homepage_started = True
        st.rerun()

    st.html(
        """
        <div class="kh-note">
            이 서비스는 사용자가 보험 내용을 이해하도록 돕는 참고 도구입니다.
            보험회사의 지급 결정, 법률상담 또는 손해사정을 대신하지 않으며
            합의금을 산정하지 않습니다.
        </div>
        """
    )

    st.stop()


# =========================================================
# 화면
# =========================================================

st.title(
    "기칸의 자동차보험 보상 나침반"
)

st.caption(
    "사고내용과 실제 가입담보를 확인해 "
    "현재 입력값 기준 예상 보험금과 다음에 준비할 자료를 보여드립니다."
)

st.subheader(
    "1. 사고내용 말하기"
)

st.caption(
    "마이크를 누르고 사고내용을 말한 뒤 종료하세요. "
    "변환된 문장은 바로 아래에서 수정할 수 있습니다."
)

recorded_audio = st.audio_input(
    "녹음 시작",
    key="accident_audio",
    label_visibility="visible",
)

if recorded_audio is not None:
    process_recorded_audio(
        recorded_audio
    )

if st.session_state.get(
    "stt_error"
):
    st.error(
        "음성을 문자로 바꾸지 못했습니다. "
        f"{st.session_state.stt_error}"
    )

    if st.button(
        "같은 녹음 다시 변환",
        key="retry_stt",
    ):
        st.session_state.force_stt_retry = True
        st.session_state.failed_audio_hash = ""
        st.session_state.stt_error = ""
        st.rerun()

elif st.session_state.get(
    "processed_audio_hash"
):
    st.success(
        "문자 변환이 완료되었습니다. "
        "아래 내용을 확인하고 필요한 부분을 수정하세요."
    )

accident_input = st.text_area(
    "사고 내용 및 수정사항",
    key="accident_input_text",
    height=190,
    placeholder=(
        "예: 2026년 9월 24일 오후 2시, "
        "교차로에서 직진 중 상대 차량과 충돌했습니다.\n"
        "부상 여부, 차량 손해, 보험회사가 제시한 과실률도 "
        "아는 범위에서 적어 주세요."
    ),
)

st.subheader(
    "2. 보험증권 촬영 및 자동입력"
)

st.caption(
    "보험증권을 촬영하면 보험회사·보험상품명·보험종류·보험기간·가입담보·가입한도를 읽어 아래 입력칸에 자동으로 채웁니다. "
    "개인정보는 보상분석에 필요하지 않은 항목을 가려서 처리하며, 자동입력 결과는 반드시 한 번 확인해 주세요."
)

policy_upload = st.file_uploader(
    "이미 찍어 둔 보험증권 사진이 있으면 선택",
    type=["jpg", "jpeg", "png", "webp", "heic", "heif"],
    key="policy_image_upload",
)

policy_source = policy_upload

if policy_source is not None:
    if st.button(
        "보험증권 읽어서 자동입력",
        key="read_policy_image_button",
        use_container_width=True,
    ):
        with st.spinner("보험증권의 가입정보를 읽고 있습니다..."):
            if process_policy_image_upload(policy_source):
                st.success("보험증권을 읽었습니다. 아래 자동입력 내용을 확인해 주세요.")
                st.rerun()

if st.session_state.get("policy_ocr_error"):
    st.error(
        "보험증권을 읽지 못했습니다. "
        + st.session_state.get("policy_ocr_error", "")
    )
ocr_result = st.session_state.get("policy_ocr_result")
if ocr_result:
    basic = ocr_result.get("basic_info") or {}
    st.success("보험증권 자동입력 결과가 준비되었습니다. 아래에서 잘못 읽힌 부분만 수정하세요.")
    summary_rows = [
        {"항목": "보험회사", "자동 판독": basic.get("보험회사") or "확인 필요"},
        {"항목": "보험상품명", "자동 판독": basic.get("보험상품명") or "확인 필요"},
        {"항목": "보험종류", "자동 판독": basic.get("보험종류") or "확인 필요"},
        {"항목": "보험기간", "자동 판독": basic.get("보험기간") or "확인 필요"},
    ]
    render_white_table(["항목", "자동 판독"], summary_rows)
    if ocr_result.get("document_review_required"):
        st.caption("사진에서 흐리거나 표준 담보와 정확히 연결되지 않은 항목은 아래에서 직접 확인해 주세요.")

st.subheader("3. 나의 입장")

# 보험종류는 보험증권 OCR 결과를 내부적으로 사용한다.
# 사용자가 별도로 선택하는 화면은 표시하지 않는다.
selected_insurance_type = st.session_state.get(
    "insurance_type_select",
    INSURANCE_TYPE_OPTIONS[0],
)

selected_position = st.radio(
    "나는 어떤 입장입니까?",
    USER_POSITIONS,
    index=0,
)

st.subheader(
    "4. 가입내용 확인"
)

policy_name_by_type = {
    "개인용 자동차보험": "개인용 약관.pdf",
    "업무용 자동차보험": "업무용 약관.pdf",
    "영업용 자동차보험": "영업용 약관.pdf",
    "이륜차 자동차보험": "이륜차 약관.pdf",
}
st.caption(
    f"현재 적용 보험종류: {selected_insurance_type} · "
    f"계산 및 약관 검색: {policy_name_by_type.get(selected_insurance_type, '확인 필요')}"
)

st.caption(
    "보험증권 사진을 읽은 경우 아래 항목이 자동으로 채워집니다. "
    "사용자는 잘못 읽힌 부분만 수정하면 됩니다."
)

storage_key = make_storage_key(
    selected_insurance_type
)

insurer = st.text_input(
    "보험회사",
    key=f"{storage_key}_insurer_v2_20260925_001",
    placeholder="예: 현대해상, 삼성화재",
    autocomplete="off",
)

insurance_product = st.text_input(
    "보험 상품명",
    key=f"{storage_key}_product_20260928_001",
    placeholder="보험증권에서 자동입력되거나 직접 입력",
    autocomplete="off",
)

insurance_period = st.text_input(
    "보험기간",
    key=f"{storage_key}_period_20260925_001",
    placeholder=(
        "예: 2026-01-01 ~ 2026-12-31"
    ),
    autocomplete="off",
)

st.markdown(
    "#### 가입담보"
)

st.caption(
    "사진에서 읽은 내용입니다. 증권과 다른 부분만 표 안에서 고쳐 주세요."
)

ocr_source_coverages = st.session_state.get(
    "policy_ocr_source_coverages",
    [],
)

editor_seed_key = (
    f"policy_editor_seed_{selected_insurance_type}"
)

current_image_hash = st.session_state.get(
    "policy_ocr_image_hash",
    "",
)

seed_hash_key = (
    f"{editor_seed_key}_image_hash"
)

# 새로운 보험증권을 읽었을 때만
# 편집표의 초기값을 OCR 결과로 다시 만든다.
if (
    current_image_hash
    and st.session_state.get(seed_hash_key)
    != current_image_hash
):
    editor_rows = []

    for item in ocr_source_coverages:
        if not isinstance(item, dict):
            continue

        original_name = str(
            item.get("original_coverage_name")
            or ""
        ).strip()

        if not original_name:
            continue

        standard_name = normalize_ocr_coverage_name(
            item.get("standard_coverage_name")
            or original_name
        )

        editor_rows.append(
            {
                "선택": True,
                "보험증권 담보명": original_name,
                "증권상 가입금액·보상한도": str(
                    item.get("original_limit")
                    or ""
                ).strip(),
                "증권상 자기부담금": str(
                    item.get("deductible")
                    or ""
                ).strip(),
                "특약·조건": str(
                    item.get("special_conditions")
                    or ""
                ).strip(),
                "앱 연결 담보": (
                    standard_name
                    if standard_name
                    in COVERAGES_BY_TYPE.get(
                        selected_insurance_type,
                        [],
                    )
                    else ""
                ),
                "확인상태": (
                    "확인 필요"
                    if item.get("needs_review")
                    else "자동 판독"
                ),
            }
        )

    st.session_state[editor_seed_key] = (
        editor_rows
    )
    st.session_state[seed_hash_key] = (
        current_image_hash
    )

# 사진을 사용하지 않았을 때는
# 빈 표에서 직접 입력할 수 있도록 한다.
if editor_seed_key not in st.session_state:
    st.session_state[editor_seed_key] = []

editor_rows = st.session_state.get(
    editor_seed_key,
    [],
)

edited_policy_rows = st.data_editor(
    editor_rows,
    key=(
        f"policy_excel_editor_"
        f"{selected_insurance_type}"
    ),
    num_rows="dynamic",
    use_container_width=True,
    hide_index=True,
    column_order=[
        "선택",
        "보험증권 담보명",
        "증권상 가입금액·보상한도",
        "증권상 자기부담금",
        "특약·조건",
    ],
    column_config={
        "선택": st.column_config.CheckboxColumn(
            "선택",
            help=(
                "보상분석에 사용할 담보만 "
                "체크하세요."
            ),
            default=True,
        ),
        "보험증권 담보명": (
            st.column_config.TextColumn(
                "보험증권 담보명",
                help=(
                    "보험증권에 적힌 담보명을 "
                    "그대로 확인·수정합니다."
                ),
                width="medium",
            )
        ),
        "증권상 가입금액·보상한도": (
            st.column_config.TextColumn(
                "증권상 가입금액·보상한도",
                help=(
                    "무한, 법정한도 등도 "
                    "증권 표현 그대로 유지합니다."
                ),
                width="medium",
            )
        ),
        "증권상 자기부담금": (
            st.column_config.TextColumn(
                "증권상 자기부담금",
                width="medium",
            )
        ),
        "특약·조건": (
            st.column_config.TextColumn(
                "특약·조건",
                width="large",
            )
        ),
        "앱 연결 담보": (
            st.column_config.TextColumn(
                "앱 연결 담보",
                help=(
                    "기존 보상분석 규칙과 연결되는 "
                    "내부 담보명입니다. 연결되지 않은 "
                    "담보는 빈칸으로 둡니다."
                ),
                width="medium",
            )
        ),
        "확인상태": (
            st.column_config.TextColumn(
                "확인상태",
                width="small",
            )
        ),
    },
)

# 사용자가 표에서 수정한 현재 값을
# 다음 rerun에서도 유지한다.
if hasattr(
    edited_policy_rows,
    "to_dict",
):
    edited_rows_list = (
        edited_policy_rows.to_dict(
            orient="records"
        )
    )
else:
    edited_rows_list = list(
        edited_policy_rows or []
    )

st.session_state[editor_seed_key] = (
    edited_rows_list
)

selected_coverages = []

for row in edited_rows_list:
    if not row.get("선택", True):
        continue

    original_name = str(
        row.get("보험증권 담보명")
        or ""
    ).strip()

    if not original_name:
        continue

    source_limit = str(
        row.get(
            "증권상 가입금액·보상한도"
        )
        or ""
    ).strip()

    source_deductible = str(
        row.get("증권상 자기부담금")
        or ""
    ).strip()

    special_conditions = str(
        row.get("특약·조건")
        or ""
    ).strip()

    linked_name = str(
        row.get("앱 연결 담보")
        or ""
    ).strip()

    if not linked_name:
        candidate = normalize_ocr_coverage_name(
            original_name
        )

        if candidate in COVERAGES_BY_TYPE.get(
            selected_insurance_type,
            [],
        ):
            linked_name = candidate

    limit_won = parse_amount_to_won(
        source_limit
    )

    deductible_won = parse_amount_to_won(
        source_deductible
    )

    # 기존 보상분석은 표에서 실제로 확인된
    # 원문을 바탕으로 만들어진 이 데이터만 사용한다.
    selected_coverages.append(
        {
            "담보명": (
                linked_name
                or original_name
            ),
            "증권 표기 담보명": (
                original_name
            ),
            "증권 원문 가입금액·보상한도": (
                source_limit
            ),
            "가입금액·보상한도": (
                str(limit_won)
                if limit_won is not None
                else source_limit
            ),
            "증권 원문 자기부담금": (
                source_deductible
            ),
            "자기부담금": (
                str(deductible_won)
                if deductible_won is not None
                else source_deductible
            ),
            "특약/비고": (
                special_conditions
            ),
            "출처 상태": (
                "보험증권 확인·수정표"
            ),
            "보상분석 연결 여부": bool(
                linked_name
            ),
        }
    )

selected_names = {
    item["담보명"]
    for item in selected_coverages
}

if (
    "자기신체사고" in selected_names
    and "자동차상해" in selected_names
):
    st.warning(
        "자기신체사고와 자동차상해가 모두 있습니다. "
        "보험증권 원문을 다시 확인해 주세요."
    )

unlinked_names = [
    item.get("증권 표기 담보명")
    for item in selected_coverages
    if not item.get(
        "보상분석 연결 여부"
    )
]

if unlinked_names:
    st.info(
        "현재 약관 분석 규칙과 자동 연결되지 않은 담보: "
        + ", ".join(unlinked_names)
        + " · 이 담보들은 다른 담보로 임의 변환하지 않고 "
        "증권 원문을 그대로 보존합니다."
    )

st.subheader("5. 예상 보험금 계산값 입력")

st.caption(
    "부상급수는 사고 당시의 상해 급수이고, 후유장애급수는 치료 후 남은 장애에 관한 급수로 서로 다릅니다. "
    "진단서나 보험회사에서 확인된 급수를 사용자가 직접 선택해야 합니다. "
    "앱이 사고 설명만 보고 부상급수나 후유장애급수를 임의로 판정하지 않습니다."
)

injury_grade = st.selectbox(
    "부상급수",
    INJURY_GRADE_OPTIONS,
    index=1,
)

disability_grade = st.selectbox(
    "후유장애급수",
    INJURY_GRADE_OPTIONS,
    index=1,
)

hospital_days = st.number_input(
    "입원일수",
    min_value=0,
    value=0,
    step=1,
    format="%d",
    help="입원하지 않았다면 0일로 입력하세요. 음수와 소수는 입력할 수 없습니다.",
)

with st.expander("보험금에 영향을 줄 수 있는 조건 확인", expanded=False):
    st.caption("확인되지 않은 값은 그대로 두고, 알고 있는 내용만 선택하세요.")

    preconditions = {}
    precondition_options = {
        "상대 차량 보험가입 상태": ["의무보험만 가입", "대인배상Ⅱ까지 가입", "무보험", "확인되지 않음", "상대 차량 없음"],
        "무보험자동차상해 적용 검토": ["해당 없음", "적용 검토 필요", "확인되지 않음"],
        "사고 운전자 면허 상태": ["유효한 면허", "무면허", "면허정지 중", "면허취소 상태", "확인되지 않음"],
        "음주운전 여부": ["아니요", "예", "확인되지 않음"],
        "마약·약물운전 여부": ["아니요", "예", "확인되지 않음"],
        "사고발생 후 조치의무 위반 여부": ["아니요", "예", "확인되지 않음"],
        "운전자 범위 한정특약 위반 여부": ["아니요", "예", "확인되지 않음"],
        "운전자 연령 한정특약 위반 여부": ["아니요", "예", "확인되지 않음"],
        "피보험자동차 해당 여부": ["예", "아니요", "확인되지 않음"],
        "차량 사용에 대한 피보험자의 허락 여부": ["허락받음", "허락받지 않음", "확인되지 않음", "해당 없음"],
        "보험증권상 차량 용도와 실제 사용 목적 일치 여부": ["일치함", "일치하지 않음", "확인되지 않음"],
        "고의사고 여부": ["아니요", "예", "확인되지 않음"],
        "안전벨트 또는 안전모 착용 여부": ["착용", "미착용", "확인되지 않음", "해당 없음"],
    }

    inferred_preconditions = infer_accident_preconditions(accident_input)
    if inferred_preconditions:
        inferred_text = ", ".join(
            f"{label}: {value}"
            for label, value in inferred_preconditions.items()
        )
        st.caption(
            "사고내용에서 자동으로 확인한 항목: " + inferred_text
            + " · 잘못 이해한 항목은 아래에서 수정하세요."
        )

    columns = st.columns(2)
    for idx, (label, options) in enumerate(precondition_options.items()):
        widget_key = f"precondition_{idx}_v1"
        if widget_key not in st.session_state:
            inferred_value = inferred_preconditions.get(label)
            if inferred_value in options:
                st.session_state[widget_key] = inferred_value
            elif "확인되지 않음" in options:
                st.session_state[widget_key] = "확인되지 않음"
            else:
                st.session_state[widget_key] = options[0]
        with columns[idx % 2]:
            value = st.selectbox(
                label,
                options,
                key=widget_key,
                help=f"{label}에 대한 확인 상태입니다.",
            )
            preconditions[label] = value

    if "무보험자동차에의한상해" in selected_names:
        st.markdown("##### 무보험자동차에 의한 상해 약관 계산 항목")
        st.caption(
            "개인용 자동차보험 약관 제20조의 계산식과 공제항목입니다. "
            "확인된 금액만 입력하면 현재 입력값 기준 예상액에 반영됩니다. "
            "빈칸은 0원으로 확정하지 않고, 아직 반영하지 않은 추가 확인 변수로 안내합니다."
        )
        uninsured_policy_fields = [
            (
                "약관상 비용(손해방지·경감비용 및 권리보전·행사비용)",
                "uninsured_policy_costs_v1",
            ),
            (
                "공제 1. 대인배상Ⅰ·책임공제·정부보장사업에서 지급될 수 있는 금액",
                "uninsured_deduction_d1_v1",
            ),
            (
                "공제 2. 배상의무자 차량의 대인배상Ⅱ·공제계약에서 지급될 수 있는 금액",
                "uninsured_deduction_obligor_d2_v1",
            ),
            (
                "공제 3. 탑승 차량의 대인배상Ⅱ·공제계약에서 지급될 수 있는 금액",
                "uninsured_deduction_occupied_d2_v1",
            ),
            (
                "공제 4. 배상의무자에게 이미 받은 손해배상금",
                "uninsured_deduction_received_obligor_v1",
            ),
            (
                "공제 5. 제3자가 부담할 금액 중 이미 받은 금액",
                "uninsured_deduction_received_third_party_v1",
            ),
        ]
        for policy_label, policy_key in uninsured_policy_fields:
            preconditions[policy_label] = render_money_input(
                policy_label,
                policy_key,
                help_text=(
                    "약관 또는 보험회사 확인자료에 있는 금액만 입력하세요. "
                    "모르면 빈칸으로 두세요."
                ),
            )

    preconditions["사고일"] = "사고 당시 적용 약관을 판단하는 기준일로 사용"


st.markdown("### 사용자 측 적용 과실비율")
user_fault_rate = st.selectbox(
    "사용자 측 적용 과실비율",
    ["아직 정해지지 않음"] + [f"{value}%" for value in range(0, 101, 10)],
    index=0,
)
if user_fault_rate != "아직 정해지지 않음":
    user_fault_value = int(user_fault_rate.replace('%', ''))
    if preconditions.get("상대 차량 보험가입 상태") != "상대 차량 없음":
        opponent_fault_value = max(0, 100 - user_fault_value)
        st.caption(f"상대방 측 과실비율: {opponent_fault_value}%")
else:
    opponent_fault_value = None

fault_source = st.radio(
    "과실비율 출처",
    [
        "아직 정해지지 않음",
        "양측 보험회사 협의",
        "자동차사고 과실비율 분쟁심의위원회",
        "법원 판결 또는 조정",
        "당사자 합의",
        "사용자가 임시로 입력한 값",
    ],
    index=0,
)

fault_rate = f"{user_fault_rate} · {fault_source}"
if user_fault_rate == "아직 정해지지 않음":
    fault_rate = f"{user_fault_rate} · {fault_source}"
else:
    fault_rate = f"{user_fault_rate} · {fault_source}"

render_actual_cost_inputs(
    selected_coverages
)

submitted = st.button(
    "예상 보험금 계산하기",
    key="confirm_inputs",
)

if submitted:
    st.session_state.analysis_result = (
        build_analysis_summary(
            position=selected_position,
            accident_text=accident_input,
            insurance_type=(
                selected_insurance_type
            ),
            insurer=insurer,
            insurance_product=insurance_product,
            insurance_period=(
                insurance_period
            ),
            injury_grade=injury_grade,
            disability_grade=disability_grade,
            hospital_days=hospital_days,
            fault_rate=fault_rate,
            selected_coverages=(
                selected_coverages
            ),
            preconditions=preconditions,
        )
    )
    result = st.session_state["analysis_result"]
    render_input_review_tables(result)
    render_policy_analysis_result(result)



# =========================================================
# 입력 결과
# =========================================================

result = st.session_state.get(
    "analysis_result"
)

if result and not submitted:
    render_input_review_tables(result)
    render_policy_analysis_result(result)


st.divider()

st.caption(
    "이 서비스는 사용자가 제공한 보험증권, 사고정보와 "
    "자동차보험 약관을 바탕으로 약관상 예상 보험금을 "
    "정리하기 위한 참고 도구입니다. "
    "실제 지급 여부와 금액은 보험회사의 사고 조사, "
    "확정된 과실률, 실제 보험계약, 제출서류와 "
    "사고일 당시 약관 및 법령에 따라 달라질 수 있습니다. "
    "이 서비스는 법률상담, 손해사정, "
    "국가기관의 유권해석 또는 보험회사의 지급결정을 "
    "대신하지 않으며 합의금을 산정하지 않습니다."
)
