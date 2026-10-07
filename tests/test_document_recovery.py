import ast
import base64
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from PIL import Image, ImageDraw, ImageFont
from pillow_heif import register_heif_opener
from pypdf import PdfWriter
from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
from openpyxl import load_workbook
from insurance_ocr.image_reader import build_document_content, read_policy_document
from insurance_ocr.validator import validate_policy_result
from insurance_ocr.compensation import estimate_partial_vehicle, verified_vehicle_rule

ROOT = Path(__file__).resolve().parents[1]


def samples():
    image = Image.new("RGB", (1400, 900), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/AppleSDGothicNeo.ttc", 45)
    except OSError:
        font = ImageFont.load_default(size=35)
    lines = ["보험가입증명서", "보험회사: 현대해상", "보험종류: 개인용 자동차보험",
             "보험기간: 2026-01-01 ~ 2027-01-01", "자기차량손해 가입금액: 2,000만원",
             "자기부담금: 20만원", "계약자: 테스트 고객"]
    if not Path("/System/Library/Fonts/AppleSDGothicNeo.ttc").exists():
        lines = ["Insurance Certificate", "Insurer: Hyundai Marine", "Personal Auto Insurance", "Period: 2026-01-01 to 2027-01-01", "Own Vehicle Damage Limit: KRW 20,000,000", "Deductible: KRW 200,000"]
    for i, line in enumerate(lines):
        draw.text((60, 60+i*100), line, font=font, fill="black")
    result = {}
    for name, fmt in [("policy.jpg", "JPEG"), ("policy.png", "PNG"), ("policy.heic", "HEIF"), ("scan.pdf", "PDF")]:
        register_heif_opener()
        buf = BytesIO(); image.save(buf, format=fmt); result[name] = buf.getvalue()
    writer = PdfWriter(); page = writer.add_blank_page(1000, 700)
    font_ref = writer._add_object(DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")}))
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"):font_ref})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 22 Tf 30 620 Td (Insurance Certificate - Hyundai Marine) Tj 0 -60 Td (Personal Auto Insurance) Tj 0 -60 Td (Period: 2026-01-01 to 2027-01-01) Tj 0 -60 Td (Own Vehicle Damage Limit: KRW 20,000,000) Tj 0 -60 Td (Deductible: KRW 200,000) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    buf=BytesIO(); writer.write(buf); result["text.pdf"] = buf.getvalue()
    return result


def app_functions():
    tree = ast.parse((ROOT/"app.py").read_text())
    functions = [n for n in tree.body if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))]
    import json, re, hashlib, os
    from html import escape
    from insurance_ocr.policy_parser import parse_amount_to_won
    state = {}
    scope = dict(st=SimpleNamespace(session_state=state), PROJECT_ROOT=ROOT,
                 POLICY_RULES=json.loads((ROOT/"data/processed/policy_rules.json").read_text()),
                 COVERAGES_BY_TYPE={"개인용 자동차보험":["대인배상Ⅰ","대인배상Ⅱ","대물배상","자기신체사고","무보험자동차에의한상해","자기차량손해"]},
                 parse_amount_to_won=parse_amount_to_won, estimate_partial_vehicle=estimate_partial_vehicle,
                 verified_vehicle_rule=verified_vehicle_rule, re=re, json=json, hashlib=hashlib, os=os, escape=escape)
    exec(compile(ast.Module(body=functions, type_ignores=[]), "app-functions", "exec"), scope)
    return scope


class RecoveryTests(unittest.TestCase):
    def test_formats_and_malformed_files(self):
        for name, data in samples().items():
            content=build_document_content(data, name)
            self.assertEqual(content["type"], "input_file" if name.endswith("pdf") else "input_image")
            if not name.endswith("pdf"):
                normalized=base64.b64decode(content["image_url"].split(",")[1])
                self.assertEqual(Image.open(BytesIO(normalized)).format, "JPEG")
        for name in ["bad.pdf", "bad.heic", "bad.jpg", "bad.txt"]:
            with self.assertRaises(ValueError):
                build_document_content(b"invalid", name)
        with self.assertRaises(ValueError):
            build_document_content(b"x"*(20*1024*1024+1), "big.pdf")
        w=PdfWriter();w.add_blank_page(100,100);w.encrypt("password");b=BytesIO();w.write(b)
        with self.assertRaises(ValueError):build_document_content(b.getvalue(),"encrypted.pdf")

    def test_api_contract_and_incomplete(self):
        from json import dumps
        output={"basic_info":{},"coverages":[],"document_review_required":True}
        fake=SimpleNamespace(status="completed",output_text=dumps(output))
        with patch("openai.OpenAI") as client:
            create=client.return_value.__enter__.return_value.responses.create
            create.return_value=fake
            for name,data in samples().items():
                self.assertEqual(read_policy_document(data,name,"fake-test-key"),dict(output,review_notes=[]))
                self.assertFalse(create.call_args.kwargs["store"])
            create.return_value=SimpleNamespace(status="incomplete")
            with self.assertRaises(RuntimeError):read_policy_document(samples()["policy.jpg"],"policy.jpg","fake-test-key")

    def test_uncertain_coverage(self):
        doc=validate_policy_result({"coverages":[{"original_coverage_name":"自己","original_limit":"","confidence":0.2}]})
        self.assertTrue(doc["coverages"][0]["needs_review"])
        self.assertIsNone(doc["coverages"][0]["limit_amount_won"])
        english=validate_policy_result({"coverages":[{"original_coverage_name":"Own Vehicle Damage Limit","original_limit":"KRW 20,000,000","deductible":"KRW 200,000","confidence":0.9}]})["coverages"][0]
        self.assertEqual(english["standard_coverage_name"],"자기차량손해")
        self.assertEqual(english["limit_amount_won"],20000000)
        self.assertFalse(english["needs_review"])

    def test_calculation_and_excel(self):
        self.assertEqual(estimate_partial_vehicle(1000000,100000,200000,20000000,15000000)[0],900000)
        self.assertEqual(estimate_partial_vehicle(0,0,0,20000000,15000000)[0],0)
        self.assertIsNone(estimate_partial_vehicle(20000000,0,200000,20000000,15000000)[0])
        self.assertIsNone(estimate_partial_vehicle(None,0,0,20,15)[0])
        scope=app_functions()
        scope["sum_relevant_costs_for_coverage"]=lambda name:1000000
        result={"insurance_type":"개인용 자동차보험","insurer":"현대해상","selected_coverages":[{"담보명":"자기차량손해","가입금액·보상한도":"2,000만원","자기부담금":"20만원"}],
                "vehicle_terms":{"damage":1000000,"expenses":100000,"vehicle_value":15000000,"confirmed":True},
                "preconditions":{"다른 보험이나 상대 보험사에서 이미 받은 금액":"0원"},"accident_text":"=HYPERLINK(\"bad\")"}
        rows,excluded=scope["build_policy_analysis_rows"](result)
        self.assertEqual(len(rows),6);self.assertEqual(len(excluded),5)
        table,total=scope["build_analysis_table"](rows);self.assertEqual(total,900000)
        self.assertEqual(scope["st"].session_state["policy_analysis_payload"]["total_estimate"],total)
        workbook=load_workbook(BytesIO(scope["build_analysis_excel"](result,rows,excluded,total)))
        self.assertEqual(workbook["보상분석표"].cell(8,9).value,"900,000원")
        self.assertEqual(workbook["확인필요항목"].max_row,6)
        self.assertEqual(workbook["입력정보"].cell(10,2).data_type,"s")
        result["vehicle_terms"]["confirmed"]=False
        rows,excluded=scope["build_policy_analysis_rows"](result)
        self.assertEqual(scope["build_analysis_table"](rows)[1],900000)
        self.assertTrue(rows[0]["조건부추정"])
        self.assertIn("총계에 포함",rows[0]["적용 약관 규정 및 산정근거"])
        self.assertIn("일부 항목 미정",scope["build_analysis_table"](rows)[0][-1]["적용 약관 규정 및 산정근거"])
        book=load_workbook(BytesIO(scope["build_analysis_excel"](result,rows,excluded,900000)))
        self.assertEqual(book["확인필요항목"].cell(2,6).value,"조건부 추정액 포함")
        self.assertEqual(book["보상분석표"].cell(8,9).value,"900,000원")
        result["vehicle_terms"]["damage"]=None
        rows,_=scope["build_policy_analysis_rows"](result)
        self.assertIsNone(rows[0]["계산가능금액"])
        self.assertEqual(scope["build_analysis_table"](rows)[1],0)
        self.assertEqual(scope["parse_optional_amount"]("사망 1억 / 부상 3천만원"),None)
        self.assertEqual(scope["parse_optional_amount"]("1.25억원"),125000000)
        for bad in ["1..2억원", "-100원", "20% 최저20만원 최고50만원", "0.00001만원"]:
            self.assertIsNone(scope["parse_optional_amount"](bad))

    def test_upload_ui_autofill_edit_and_blank(self):
        from streamlit.testing.v1 import AppTest
        import streamlit
        doc={"basic_info":{"보험회사":"현대해상","보험기간":"2026-01-01 ~ 2027-01-01"},"coverages":[{"original_coverage_name":"자기차량손해","original_limit":"2,000만원","deductible":"20만원","confidence":0.99}]}
        upload=SimpleNamespace(name="policy.jpg",getvalue=lambda:b"mock-image")
        with patch.object(streamlit,"file_uploader",return_value=upload), patch("insurance_ocr.image_reader.read_policy_document",return_value=doc):
            app=AppTest.from_file(str(ROOT/"app.py"),default_timeout=30).run()
            self.assertEqual(len(app.exception),0)
            app.button(key="homepage_start_button").click().run()
            self.assertEqual(len(app.exception),0)
            app.button(key="read_policy_document").click().run()
            self.assertEqual(len(app.exception),0)
            scope=app_functions()
            insurer_key="개인용_자동차보험_insurer_v2_20260925_001"
            self.assertEqual(app.text_input(key=insurer_key).value,"현대해상")
            app.text_input(key=insurer_key).set_value("수정 보험사").run()
            self.assertEqual(app.text_input(key=insurer_key).value,"수정 보험사")
            manual_input=next(x for x in app.text_input if x.label=="자기신체사고 임시 추정액 (원)")
            manual_input.set_value("500000").run()
            app.button(key="confirm_inputs").click().run()
            self.assertEqual(len(app.exception),0)
            self.assertEqual(len(app.session_state["policy_analysis_payload"]["rows"]),6)
            self.assertEqual(app.session_state["policy_analysis_payload"]["total_estimate"],500000)
            metrics={m.label:m.value for m in app.metric}
            self.assertEqual(metrics["추정 보험금 총계"],"500,000원")
            self.assertEqual(metrics["총계 중 사용자 임시 입력액"],"500,000원")
            self.assertEqual(metrics["금액 미정 담보"],"5건")
            doc["basic_info"]={}
            doc["coverages"][0]["confidence"]=0.2
            app.button(key="read_policy_document").click().run()
            self.assertEqual(len(app.exception),0)
            self.assertEqual(app.text_input(key=insurer_key).value,"")
            amount_inputs=[x for x in app.text_input if x.label=="가입금액·보상한도"]
            self.assertEqual(amount_inputs[0].value,"")

    def test_manual_estimates_in_screen_and_excel(self):
        scope=app_functions()
        scope["sum_relevant_costs_for_coverage"]=lambda name:None
        result={"insurance_type":"개인용 자동차보험", "selected_coverages":[],
                "manual_estimates":{"자기신체사고":{"amount":500000,"reason":"급별 한도 확인 필요"},
                                    "무보험자동차에의한상해":{"amount":300000,"reason":"중복 지급 조정 확인 필요"},
                                    "대물배상":{"amount":0,"reason":"손해 없음"}}}
        rows,review=scope["build_policy_analysis_rows"](result)
        table,total=scope["build_analysis_table"](rows)
        self.assertEqual(total,800000)
        self.assertEqual(scope["st"].session_state["policy_analysis_payload"]["total_estimate"],800000)
        self.assertEqual(sum(bool(r.get("사용자추정")) for r in rows),3)
        self.assertIn("중복 지급 조정", next(r for r in rows if r["계산가능금액"]==300000)["적용 약관 규정 및 산정근거"])
        workbook=load_workbook(BytesIO(scope["build_analysis_excel"](result,rows,review,total)))
        self.assertEqual(workbook["보상분석표"].cell(8,9).value,"800,000원")
        states=[r[5] for r in workbook["확인필요항목"].iter_rows(min_row=2,values_only=True)]
        self.assertEqual(states.count("사용자 임시 추정액 포함"),3)
        result["manual_estimates"]={}
        rows,_=scope["build_policy_analysis_rows"](result)
        self.assertTrue(all(r["계산가능금액"] is None for r in rows))

    def test_voice_unchanged(self):
        import hashlib, json
        expected=json.loads((ROOT/"tests/voice_function_hashes.json").read_text())
        functions={n.name:n for n in ast.parse((ROOT/"app.py").read_text()).body if isinstance(n,ast.FunctionDef)}
        for name,fingerprint in expected.items():
            self.assertEqual(hashlib.sha256(ast.dump(functions[name]).encode()).hexdigest(),fingerprint,name)

if __name__ == "__main__":unittest.main()
